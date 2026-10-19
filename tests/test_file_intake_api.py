import hashlib
import io
import time
from pathlib import Path

import pytest
import xlwt
from fastapi.testclient import TestClient
from openpyxl import Workbook
from reportlab.pdfgen.canvas import Canvas

from project_watch.api import create_app
from project_watch.extractors import ExtractionError


def _project(client: TestClient) -> str:
    response = client.post("/api/projects", json={"name": "Synthetic Company"})
    assert response.status_code == 201
    return response.json()["id"]


def test_upload_is_queued_then_worker_exposes_located_text(tmp_path: Path) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)
    payload = b"Synthetic company profile\nSynthetic product"

    response = client.post(
        f"/api/projects/{project_id}/files",
        files={"file": ("intro.txt", payload, "text/plain")},
    )

    assert response.status_code == 202
    intake = response.json()
    assert intake["source"]["content_sha256"] == hashlib.sha256(payload).hexdigest()
    assert intake["job"]["status"] == "queued"
    assert client.get(f"/api/projects/{project_id}/intake").json()[0]["status"] == "queued"

    processed = app.state.intake_worker.process_one()

    assert processed.status == "completed"
    fragments = client.get(f"/api/sources/{intake['source']['id']}/fragments").json()
    assert [(row["locator"], row["text"]) for row in fragments] == [
        ("line 1", "Synthetic company profile"),
        ("line 2", "Synthetic product"),
    ]
    assert client.get(f"/api/projects/{project_id}/summary").json()["confirmed_count"] == 0


def test_legacy_xls_upload_is_parsed_without_confirming_numbers(tmp_path: Path) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Synthetic")
    sheet.write(0, 0, "2025 actual")
    sheet.write(1, 0, 120)
    payload = io.BytesIO()
    workbook.save(payload)

    uploaded = client.post(
        f"/api/projects/{project_id}/files",
        files={"file": ("synthetic.xls", payload.getvalue(), "application/vnd.ms-excel")},
    )
    assert uploaded.status_code == 202
    assert app.state.intake_worker.process_one().status == "completed"
    fragments = client.get(f"/api/sources/{uploaded.json()['source']['id']}/fragments").json()
    assert "A2[header:2025 actual]=120" in fragments[1]["text"]
    assert client.get(f"/api/projects/{project_id}/summary").json()["confirmed_count"] == 0


def test_xlsx_upload_proposes_located_financial_value_then_human_confirms(tmp_path: Path) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "合并利润表"
    sheet.append(["单位：万元"])
    sheet.append(["项目", "2025年实际"])
    sheet.append(["营业收入", 120])
    payload = io.BytesIO()
    workbook.save(payload)
    uploaded = client.post(
        f"/api/projects/{project_id}/files",
        files={
            "file": (
                "synthetic.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 202
    processed = app.state.intake_worker.process_one()
    assert processed.status == "completed", processed.error_code
    observations = client.get(f"/api/projects/{project_id}/observations").json()
    finance = [item for item in observations if item["statement"] == "营业收入"]
    assert len(finance) == 1
    assert (finance[0]["number"], finance[0]["period"], finance[0]["basis"]) == (
        "120",
        "2025",
        "actual",
    )
    assert finance[0]["source_locator"] == "sheet 合并利润表 row 3"
    assert finance[0]["status"] == "candidate"
    screen = f"/api/projects/{project_id}/financial-screen"
    assert client.get(screen).json()["accepted_count"] == 0
    approved = client.post(
        f"/api/projects/{project_id}/observations/{finance[0]['id']}/confirm",
        json={"reviewer": "synthetic-reviewer", "expected_version": 1},
    )
    assert approved.status_code == 200
    assert client.get(screen).json()["accepted_count"] == 1


def test_xlsx_percent_display_enters_operating_screen_only_after_review(tmp_path: Path) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产能统计"
    sheet.append(["产品", "指标", "单位", "2024年实际", "2025年实际"])
    sheet.append(["连接器A", "产能利用率", "%", 0.6, 0.7])
    sheet["D2"].number_format = "0%"
    sheet["E2"].number_format = "0%"
    payload = io.BytesIO()
    workbook.save(payload)
    uploaded = client.post(
        f"/api/projects/{project_id}/files",
        files={
            "file": (
                "synthetic.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 202
    assert app.state.intake_worker.process_one().status == "completed"
    observations = client.get(f"/api/projects/{project_id}/observations").json()
    utilization = [item for item in observations if item["statement"] == "产能利用率"]
    assert sorted((item["number"], item["period"], item["unit"]) for item in utilization) == [
        ("60.0", "2024", "%"),
        ("70.0", "2025", "%"),
    ]
    assert all(item["status"] == "candidate" for item in utilization)
    screen = f"/api/projects/{project_id}/monitoring/domains"
    assert client.get(screen).json()["operating_metrics"]["actual_count"] == 0
    reviewed = client.post(
        f"/api/projects/{project_id}/observations/batch-review",
        json={
            "reviewer": "synthetic-reviewer",
            "items": [
                {"observation_id": item["id"], "expected_version": 1, "action": "confirm"}
                for item in utilization
            ],
        },
    )
    assert reviewed.status_code == 200
    result = client.get(screen).json()["operating_metrics"]
    assert result["actual_count"] == 2
    assert {item["period"]: item["value"] for item in result["evidence"]} == {
        "2024": "60.0",
        "2025": "70.0",
    }


def test_text_formatted_financial_amounts_remain_pending_after_upload(tmp_path: Path) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "合并利润表"
    sheet.append(["单位：万元"])
    sheet.append(["项目", "2025年实际", "2026年预测"])
    sheet.append(["净利润", "(1,234.50)", "2,000.00"])
    payload = io.BytesIO()
    workbook.save(payload)

    uploaded = client.post(
        f"/api/projects/{project_id}/files",
        files={
            "file": (
                "synthetic.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 202
    assert app.state.intake_worker.process_one().status == "completed"

    observations = client.get(f"/api/projects/{project_id}/observations").json()
    finance = [item for item in observations if item["statement"] == "净利润"]
    assert sorted((item["number"], item["period"], item["basis"]) for item in finance) == [
        ("-1234.50", "2025", "actual"),
        ("2000.00", "2026", "forecast"),
    ]
    assert all(
        item["status"] == "candidate" and item["source_locator"] == "sheet 合并利润表 row 3"
        for item in finance
    )
    assert client.get(f"/api/projects/{project_id}/financial-screen").json()["accepted_count"] == 0


def test_financial_mapping_spreads_limit_across_workbook_sheets(tmp_path: Path) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)
    workbook = Workbook()
    income = workbook.active
    income.title = "合并利润表"
    income.append(["单位：万元"])
    income.append(["项目", *(f"{year}年实际" for year in range(2000, 2030))])
    income.append(["营业收入", *range(100, 130)])
    for sheet_name, metric in (
        ("合并资产负债表", "资产总计"),
        ("合并现金流量表", "经营活动现金流量净额"),
    ):
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(["单位：万元"])
        sheet.append(["项目", "2025年实际"])
        sheet.append([metric, 500])
    payload = io.BytesIO()
    workbook.save(payload)

    uploaded = client.post(
        f"/api/projects/{project_id}/files",
        files={
            "file": (
                "synthetic-three-statements.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 202
    assert app.state.intake_worker.process_one().status == "completed"
    source_id = uploaded.json()["source"]["id"]

    mapping = client.get(f"/api/projects/{project_id}/sources/{source_id}/financial-mapping")
    assert mapping.status_code == 200
    assert {k: mapping.json()[k] for k in ("candidate_count", "truncated_count", "issues")} == {"candidate_count": 24, "truncated_count": 8, "issues": []}
    assert len(mapping.json()["review_drafts"]) == 32
    candidates = [
        row
        for row in client.get(f"/api/projects/{project_id}/observations").json()
        if row["status"] == "candidate"
        and row["statement"] in {"营业收入", "资产总计", "经营活动现金流量净额"}
    ]
    assert len(candidates) == 24
    assert {row["statement"] for row in candidates} == {
        "营业收入",
        "资产总计",
        "经营活动现金流量净额",
    }


def test_foreign_currency_sheet_requires_review_instead_of_rmb_candidate(tmp_path: Path) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "合并利润表"
    sheet.append(["币种：美元", "单位：万元"])
    sheet.append(["项目", "2025年实际"])
    sheet.append(["营业收入", 120])
    payload = io.BytesIO()
    workbook.save(payload)

    uploaded = client.post(
        f"/api/projects/{project_id}/files",
        files={
            "file": (
                "synthetic.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 202
    assert app.state.intake_worker.process_one().status == "completed"
    source_id = uploaded.json()["source"]["id"]
    mapping = client.get(f"/api/projects/{project_id}/sources/{source_id}/financial-mapping")
    assert mapping.status_code == 200
    assert mapping.json() == {
        "candidate_count": 0,
        "truncated_count": 0,
        "issues": [{"source_locator": "sheet 合并利润表 row 3", "reason": "currency_unsupported"}],
        "review_drafts": [],
        "review_truncated_count": 0,
    }
    assert not any(
        item["statement"] == "营业收入"
        for item in client.get(f"/api/projects/{project_id}/observations").json()
    )


def test_xlsx_mapping_exposes_missing_unit_without_inventing_a_number(tmp_path: Path) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "合并利润表"
    sheet.append(["项目", "2025年实际"])
    sheet.append(["营业收入", 120])
    payload = io.BytesIO()
    workbook.save(payload)
    uploaded = client.post(
        f"/api/projects/{project_id}/files",
        files={
            "file": (
                "synthetic.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    source_id = uploaded.json()["source"]["id"]
    path = f"/api/projects/{project_id}/sources/{source_id}/financial-mapping"
    assert client.get(path).status_code == 409
    assert app.state.intake_worker.process_one().status == "completed"

    response = client.get(path)
    assert response.status_code == 200
    assert {k: response.json()[k] for k in ("candidate_count", "truncated_count", "issues")} == {
        "candidate_count": 0,
        "truncated_count": 0,
        "issues": [{"source_locator": "sheet 合并利润表 row 2", "reason": "unit_missing"}],
    }
    draft = response.json()["review_drafts"][0]
    assert draft["number"] == "120" and draft["unit"] is None
    assert draft["missing_fields"] == ["unit"]
    assert not any(
        item["statement"] == "营业收入"
        for item in client.get(f"/api/projects/{project_id}/observations").json()
    )
    another_project = client.post("/api/projects", json={"name": "Other Synthetic Company"}).json()[
        "id"
    ]
    assert (
        client.get(
            f"/api/projects/{another_project}/sources/{source_id}/financial-mapping"
        ).status_code
        == 404
    )


def test_reupload_does_not_create_second_job(tmp_path: Path) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)
    files = {"file": ("intro.txt", b"synthetic", "text/plain")}

    first = client.post(f"/api/projects/{project_id}/files", files=files).json()
    second = client.post(f"/api/projects/{project_id}/files", files=files).json()

    assert second["source"]["id"] == first["source"]["id"]
    assert second["job"]["id"] == first["job"]["id"]
    assert len(client.get(f"/api/projects/{project_id}/intake").json()) == 1


def test_financial_upload_creates_only_located_system_candidates(tmp_path: Path) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)
    uploaded = client.post(
        f"/api/projects/{project_id}/files",
        files={
            "file": (
                "财务报表.txt",
                "2025 年营业收入 1,200 万元。\n净利润仍需核对。".encode(),
                "text/plain",
            )
        },
    )
    source_id = uploaded.json()["source"]["id"]

    assert app.state.intake_worker.process_one().status == "completed"
    indexed = client.get(f"/api/projects/{project_id}/sources").json()
    assert indexed[0]["topic"] == "财务表现"
    assert indexed[0]["topic_basis"] == "filename_and_content"
    candidates = client.get(f"/api/projects/{project_id}/observations").json()
    assert {item["statement"] for item in candidates} == {
        "2025 年营业收入 1,200 万元。",
        "净利润仍需核对。",
    }
    assert all(item["source_id"] == source_id for item in candidates)
    assert all(item["status"] == "candidate" for item in candidates)
    audit = client.get(f"/api/projects/{project_id}/audit").json()
    assert len([row for row in audit if row["channel"] == "native_intake"]) == 2
    assert all(row["actor_kind"] == "system" for row in audit if row["channel"] == "native_intake")
    assert client.get(f"/api/projects/{project_id}/summary").json()["confirmed_count"] == 0
    reviewed = client.post(
        f"/api/observations/{candidates[0]['id']}/confirm",
        json={"reviewer": "synthetic-analyst", "expected_version": 1},
    )
    assert reviewed.status_code == 200
    assert client.get(f"/api/projects/{project_id}/summary").json()["confirmed_count"] == 1

    changed = client.post(
        f"/api/projects/{project_id}/files",
        files={
            "file": (
                "财务报表.txt",
                "2026 年营业收入仍需核对。".encode(),
                "text/plain",
            )
        },
    )
    assert changed.json()["source"]["id"] != source_id
    assert app.state.intake_worker.process_one().status == "completed"
    history = client.get(f"/api/projects/{project_id}/observations").json()
    assert any(
        item["id"] == candidates[0]["id"] and item["status"] == "confirmed" for item in history
    )
    assert any(
        item["source_id"] == changed.json()["source"]["id"]
        and item["statement"] == "2026 年营业收入仍需核对。"
        and item["status"] == "candidate"
        for item in history
    )
    impacts = client.get(f"/api/projects/{project_id}/review-impacts").json()
    assert len(impacts) == 1
    assert impacts[0]["observation_id"] == candidates[0]["id"]
    assert impacts[0]["trigger_source_id"] == changed.json()["source"]["id"]
    assert impacts[0]["status"] == "pending"
    pending_brief = client.get(f"/api/projects/{project_id}/report/current").json()
    assert pending_brief["review_impacts"][0]["old_observation_id"] == candidates[0]["id"]
    assert pending_brief["review_impacts"][0]["new_source_id"] == changed.json()["source"]["id"]
    assert pending_brief["review_impacts"][0]["status"] == "pending"
    resolved = client.post(
        f"/api/projects/{project_id}/review-impacts/{impacts[0]['id']}/resolve",
        json={
            "resolution": "follow_up",
            "note": "新旧口径仍需财务访谈核实。",
            "reviewer": "synthetic-analyst",
        },
    )
    assert resolved.status_code == 200
    assert resolved.json()["status"] == "reviewed"
    assert resolved.json()["note"] == "新旧口径仍需财务访谈核实。"
    reviewed_brief = client.get(f"/api/projects/{project_id}/report/current").json()
    assert reviewed_brief["pending_review_count"] == 0
    assert reviewed_brief["review_impacts"][0]["resolution"] == "follow_up"
    assert reviewed_brief["review_impacts"][0]["note"] == "新旧口径仍需财务访谈核实。"
    assert client.get(f"/api/projects/{project_id}/observations").json()[0]["status"] == "confirmed"
    assert (
        client.post(
            f"/api/projects/{project_id}/review-impacts/{impacts[0]['id']}/resolve",
            json={"resolution": "unchanged", "note": "重复提交", "reviewer": "synthetic-analyst"},
        ).status_code
        == 409
    )


def test_mixed_upload_prompts_two_domains_without_changing_confirmed_facts(tmp_path: Path) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)
    confirmed_ids: set[str] = set()
    for filename, text, subject in (
        ("财务记录.txt", "营业收入仍需核对。", "财务表现"),
        ("产品记录.txt", "产品路线仍需核对。", "产品与技术"),
    ):
        uploaded = client.post(
            f"/api/projects/{project_id}/files",
            files={"file": (filename, text.encode(), "text/plain")},
        )
        assert uploaded.status_code == 202
        assert app.state.intake_worker.process_one().status == "completed"
        candidates = [
            item
            for item in client.get(f"/api/projects/{project_id}/observations").json()
            if item["source_id"] == uploaded.json()["source"]["id"]
        ]
        assert len(candidates) == 1
        assert candidates[0]["subject"] == subject
        confirmed = client.post(
            f"/api/observations/{candidates[0]['id']}/confirm",
            json={"reviewer": "synthetic-analyst", "expected_version": 1},
        )
        assert confirmed.status_code == 200
        confirmed_ids.add(candidates[0]["id"])

    changed = client.post(
        f"/api/projects/{project_id}/files",
        files={
            "file": (
                "财务更新.txt",
                "营业收入出现变化。\n产品路线出现变化。".encode(),
                "text/plain",
            )
        },
    )
    assert changed.status_code == 202
    assert app.state.intake_worker.process_one().status == "completed"
    impacts = client.get(f"/api/projects/{project_id}/review-impacts").json()
    assert len(impacts) == 2
    assert {item["observation_id"] for item in impacts} == confirmed_ids
    assert all(item["trigger_source_id"] == changed.json()["source"]["id"] for item in impacts)
    brief = client.get(f"/api/projects/{project_id}/report/current").json()
    assert brief["pending_review_count"] == 2
    assert {item["old_observation_id"] for item in brief["review_impacts"]} == confirmed_ids
    assert all(
        item["status"] == "confirmed"
        for item in client.get(f"/api/projects/{project_id}/observations").json()
        if item["id"] in confirmed_ids
    )


def test_read_only_source_comparison_requires_two_completed_project_documents(
    tmp_path: Path,
) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)
    other_id = client.post("/api/projects", json={"name": "Other Synthetic Company"}).json()["id"]

    def upload(project: str, content: bytes) -> str:
        response = client.post(
            f"/api/projects/{project}/files",
            files={"file": ("report.txt", content, "text/plain")},
        )
        assert response.status_code == 202
        return response.json()["source"]["id"]

    before = upload(project_id, b"Revenue: 10\nSame")
    after = upload(project_id, b"Revenue: 12\nSame")
    outside = upload(other_id, b"Private synthetic source")
    url = f"/api/projects/{project_id}/sources/{after}/compare"
    assert client.get(url, params={"before_source_id": before}).status_code == 409
    assert client.get(url, params={"before_source_id": outside}).status_code == 404
    assert client.get(url, params={"before_source_id": after}).status_code == 422

    for _ in range(2):
        assert app.state.intake_worker.process_one().status == "completed"
    result = client.get(url, params={"before_source_id": before})
    assert result.status_code == 200
    assert result.json()["counts"] == {"added": 0, "changed": 1, "removed": 0}
    assert result.json()["changes"][0] == {
        "locator": "line 1",
        "kind": "changed",
        "before": "Revenue: 10",
        "after": "Revenue: 12",
    }
    assert result.json()["before_source_id"] == before
    assert result.json()["after_source_id"] == after
    assert client.get(f"/api/projects/{project_id}/observations").json() == []


def test_invalid_file_or_project_does_not_create_job(tmp_path: Path) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)

    bad = client.post(
        f"/api/projects/{project_id}/files",
        files={"file": ("malware.exe", b"synthetic", "application/octet-stream")},
    )
    missing = client.post(
        "/api/projects/missing/files",
        files={"file": ("intro.txt", b"synthetic", "text/plain")},
    )

    assert bad.status_code == 422
    assert missing.status_code == 404
    assert client.get(f"/api/projects/{project_id}/intake").json() == []
    assert client.get(f"/api/projects/{project_id}/sources").json() == []


def test_worker_failure_is_a_bounded_status_and_can_retry(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)
    response = client.post(
        f"/api/projects/{project_id}/files",
        files={"file": ("intro.txt", b"Synthetic content", "text/plain")},
    )
    job_id = response.json()["job"]["id"]

    def fail_extraction(*_args, **_kwargs):
        raise ExtractionError("invalid_document")

    monkeypatch.setattr("project_watch.intake_worker.extract_document", fail_extraction)

    failed = app.state.intake_worker.process_one()

    assert failed.status == "failed"
    assert failed.error_code == "invalid_document"
    assert "Traceback" not in str(client.get(f"/api/projects/{project_id}/intake").json())
    retried = client.post(f"/api/intake/{job_id}/retry")
    assert retried.status_code == 200
    assert retried.json()["status"] == "queued"


def test_local_lifespan_worker_processes_without_manual_click(tmp_path: Path) -> None:
    app = create_app(tmp_path / "watch.sqlite3", start_worker=True)
    with TestClient(app) as client:
        project_id = _project(client)
        uploaded = client.post(
            f"/api/projects/{project_id}/files",
            files={"file": ("intro.txt", b"Synthetic introduction", "text/plain")},
        )
        assert uploaded.status_code == 202
        deadline = time.monotonic() + 3
        while time.monotonic() < deadline:
            jobs = client.get(f"/api/projects/{project_id}/intake").json()
            if jobs[0]["status"] == "completed":
                break
            time.sleep(0.05)
        assert jobs[0]["status"] == "completed"


def test_image_only_pdf_is_marked_for_ocr(tmp_path: Path) -> None:
    app = create_app(tmp_path / "watch.sqlite3")
    client = TestClient(app)
    project_id = _project(client)
    payload = io.BytesIO()
    canvas = Canvas(payload)
    canvas.rect(72, 72, 100, 100)
    canvas.save()
    uploaded = client.post(
        f"/api/projects/{project_id}/files",
        files={"file": ("scan.pdf", payload.getvalue(), "application/pdf")},
    )

    assert uploaded.status_code == 202
    result = app.state.intake_worker.process_one()
    assert result.status == "needs_ocr"
    source_id = uploaded.json()["source"]["id"]
    assert client.get(f"/api/sources/{source_id}/fragments").json() == []
