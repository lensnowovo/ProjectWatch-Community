import shutil
import subprocess
from pathlib import Path

import pytest
import xlwt
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen.canvas import Canvas

from project_watch.extractors import ExtractionError, extract_document


def test_text_lines_have_stable_locators(tmp_path: Path) -> None:
    path = tmp_path / "intro.txt"
    path.write_text("公司概况\n产品介绍\n", encoding="utf-8")

    result = extract_document(path)

    assert [(row.locator, row.text) for row in result.fragments] == [
        ("line 1", "公司概况"),
        ("line 2", "产品介绍"),
    ]
    assert result.needs_ocr is False


def test_pdf_pages_are_located(tmp_path: Path) -> None:
    path = tmp_path / "overview.pdf"
    canvas = Canvas(str(path))
    canvas.drawString(72, 700, "Synthetic market report")
    canvas.showPage()
    canvas.drawString(72, 700, "Synthetic customer data")
    canvas.save()

    result = extract_document(path)

    assert [row.locator for row in result.fragments] == ["page 1", "page 2"]
    assert "Synthetic customer data" in result.fragments[1].text
    assert result.needs_ocr is False


def test_image_only_pdf_reports_ocr_gap(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    canvas = Canvas(str(path))
    canvas.rect(72, 72, 100, 100)
    canvas.save()

    from unittest.mock import patch

    with patch("project_watch.extractors._ocr_pdf_page", return_value=None):
        result = extract_document(path)

    assert result.fragments == []
    assert result.needs_ocr is True


def test_mixed_pdf_keeps_page_locators_and_ocr_provenance(tmp_path: Path) -> None:
    from unittest.mock import patch

    path = tmp_path / "mixed.pdf"
    canvas = Canvas(str(path))
    canvas.drawString(72, 700, "Synthetic source page with enough native text for extraction")
    canvas.showPage()
    canvas.rect(72, 72, 100, 100)
    canvas.save()

    with patch("project_watch.extractors._ocr_pdf_page", return_value="Revenue 12345") as ocr:
        result = extract_document(path)

    assert [item.locator for item in result.fragments] == ["page 1", "page 2 OCR"]
    assert result.fragments[1].text == "Revenue 12345"
    assert result.needs_ocr is False
    assert ocr.call_count == 1
    assert ocr.call_args.args[:2] == (path, 2)


def test_failed_ocr_retains_partial_native_text_and_gap(tmp_path: Path) -> None:
    from unittest.mock import patch

    path = tmp_path / "partial.pdf"
    canvas = Canvas(str(path))
    canvas.drawString(72, 700, "12")
    canvas.drawImage(ImageReader(Image.new("RGB", (100, 100), "white")), 72, 100, 100, 100)
    canvas.save()

    with patch("project_watch.extractors._ocr_pdf_page", return_value=None):
        result = extract_document(path)

    assert result.fragments[0].locator == "page 1"
    assert result.needs_ocr is True


def test_local_ocr_reads_synthetic_scanned_pdf(tmp_path: Path) -> None:
    if not shutil.which("pdftoppm") or not shutil.which("tesseract"):
        pytest.skip("local OCR binaries unavailable")
    languages = subprocess.run(
        ["tesseract", "--list-langs"],
        capture_output=True,
        text=True,
        check=False,
    )
    if "chi_sim" not in languages.stdout:
        pytest.skip("Simplified Chinese OCR data unavailable")
    image = Image.new("RGB", (1200, 400), "white")
    ImageDraw.Draw(image).text(
        (80, 110),
        "TOTAL ASSETS 12345",
        fill="black",
        font=ImageFont.load_default(size=72),
    )
    path = tmp_path / "scanned.pdf"
    canvas = Canvas(str(path))
    canvas.drawImage(ImageReader(image), 72, 350, 468, 156)
    canvas.save()

    result = extract_document(path)

    assert result.needs_ocr is False
    assert result.fragments[0].locator == "page 1 OCR"
    assert "12345" in result.fragments[0].text


def test_docx_paragraphs_and_tables_are_located(tmp_path: Path) -> None:
    path = tmp_path / "brief.docx"
    document = Document()
    document.add_paragraph("Synthetic company profile")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Revenue"
    table.cell(0, 1).text = "120"
    document.save(path)

    result = extract_document(path)

    assert any(row.locator == "paragraph 1" for row in result.fragments)
    assert any(row.locator == "table 1 row 1" and "Revenue" in row.text for row in result.fragments)


def test_content_addressed_file_uses_declared_suffix(tmp_path: Path) -> None:
    path = tmp_path / "brief.docx"
    document = Document()
    document.add_paragraph("Synthetic company profile")
    document.save(path)
    stored_path = tmp_path / ("a" * 64)
    path.rename(stored_path)

    result = extract_document(stored_path, suffix=".docx")

    assert result.fragments[0].text == "Synthetic company profile"


def test_pptx_slides_are_located(tmp_path: Path) -> None:
    path = tmp_path / "bp.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text = "Synthetic product overview"
    presentation.save(path)

    result = extract_document(path)

    assert [(row.locator, row.text) for row in result.fragments] == [
        ("slide 1", "Synthetic product overview")
    ]


def test_xlsx_sheet_rows_are_located(tmp_path: Path) -> None:
    path = tmp_path / "finance.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Financials"
    sheet.append(["Revenue", "2025"])
    sheet.append([120, 140])
    workbook.save(path)

    result = extract_document(path)

    assert [row.locator for row in result.fragments] == [
        "sheet Financials row 1",
        "sheet Financials row 2",
    ]
    assert "Revenue" in result.fragments[0].text


def test_xls_sheet_rows_are_located(tmp_path: Path) -> None:
    path = tmp_path / "old-finance.xls"
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Financials")
    sheet.write(0, 0, "Revenue")
    sheet.write(0, 1, "2025 actual")
    sheet.write(1, 0, "Sales")
    sheet.write(1, 1, 120)
    workbook.save(str(path))

    result = extract_document(path)

    assert [item.locator for item in result.fragments] == [
        "sheet Financials row 1",
        "sheet Financials row 2",
    ]
    assert "B1=2025 actual" in result.fragments[0].text
    assert "B2[header:2025 actual]=120" in result.fragments[1].text


def test_mixed_actual_and_forecast_columns_keep_individual_headers(tmp_path: Path) -> None:
    path = tmp_path / "mixed.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Metric", "2025 actual", "2026 forecast"])
    sheet.append(["Sales", 120, 180])
    workbook.save(path)

    result = extract_document(path)

    row = result.fragments[1].text
    assert "B2[header:2025 actual]=120" in row
    assert "C2[header:2026 forecast]=180" in row


def test_text_formatted_financial_numbers_keep_column_headers(tmp_path: Path) -> None:
    path = tmp_path / "text-finance.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["项目", "2025年实际", "2026年预测"])
    sheet.append(["净利润", "(1,234.50)", "2,000.00"])
    workbook.save(path)

    result = extract_document(path)

    assert "B2[header:2025年实际]=(1,234.50)" in result.fragments[1].text
    assert "C2[header:2026年预测]=2,000.00" in result.fragments[1].text


def test_invalid_xls_header_fails_before_parsing(tmp_path: Path) -> None:
    path = tmp_path / "fake.xls"
    path.write_bytes(b"not an OLE file")
    with pytest.raises(ExtractionError, match="invalid_document"):
        extract_document(path)


def test_xls_beyond_1000_rows_preserves_tail(tmp_path: Path) -> None:
    path = tmp_path / "huge.xls"
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Sheet")
    sheet.write(1000, 0, "outside limit")
    workbook.save(str(path))
    result = extract_document(path)
    assert result.fragments[-1].locator == "sheet Sheet row 1001"
    assert "outside limit" in result.fragments[-1].text


def test_corrupt_office_archive_fails_explicitly(tmp_path: Path) -> None:
    path = tmp_path / "bad.docx"
    path.write_bytes(b"PK\x03\x04not-a-real-archive")

    with pytest.raises(ExtractionError) as caught:
        extract_document(path)

    assert caught.value.error_code == "invalid_document"


def test_large_sheet_preserves_all_rows(tmp_path: Path) -> None:
    path = tmp_path / "huge.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for row in range(1, 1002):
        sheet.cell(row, 1, row)
    workbook.save(path)

    result = extract_document(path)
    assert len(result.fragments) == 1001
    assert result.fragments[-1].locator == "sheet Sheet row 1001"
    assert "A1001=1001" in result.fragments[-1].text
