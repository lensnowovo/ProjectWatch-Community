"""Document size is not a semantic reason to discard valid research material."""

import io
import json
from dataclasses import asdict

from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen.canvas import Canvas

from project_watch.blob_store import BlobStore
from project_watch.extraction_process import _decode_result, extract_document_subprocess
from project_watch.extractors import _Collector, extract_document


def test_pdf_beyond_200_pages_preserves_last_page(tmp_path):
    path = tmp_path / "long.pdf"
    canvas = Canvas(str(path))
    for number in range(1, 210):
        canvas.drawString(72, 720, f"Research evidence on page {number}")
        canvas.showPage()
    canvas.save()
    result = extract_document_subprocess(path, suffix=".pdf")
    assert len(result.fragments) == 209
    assert result.fragments[-1].locator == "page 209"
    assert "209" in result.fragments[-1].text


def test_presentation_beyond_200_slides_preserves_last_slide(tmp_path):
    path = tmp_path / "long.pptx"
    presentation = Presentation()
    for number in range(1, 202):
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = f"Research {number}"
    presentation.save(path)
    result = extract_document(path)
    assert result.fragments[-1].locator == "slide 201"
    assert "201" in result.fragments[-1].text


def test_workbook_beyond_sheet_and_column_limits_preserves_tail(tmp_path):
    path = tmp_path / "wide.xlsx"
    workbook = Workbook()
    for number in range(1, 52):
        sheet = workbook.active if number == 1 else workbook.create_sheet()
        sheet.title = f"Research{number}"
        sheet.cell(1, 101, f"tail-{number}")
    workbook.save(path)
    result = extract_document(path)
    assert len(result.fragments) == 51
    assert result.fragments[-1].locator == "sheet Research51 row 1"
    assert "CW1=tail-51" in result.fragments[-1].text


def test_collector_and_receipt_preserve_over_500_fragments_and_2m_characters():
    collector = _Collector()
    for number in range(1, 602):
        collector.add(f"line {number}", str(number) + "x" * 4000)
    result = collector.finish()
    payload = json.dumps({"fragments": [asdict(f) for f in result.fragments],
                          "needs_ocr": False}).encode()
    decoded = _decode_result(payload)
    assert len(decoded.fragments) == 601
    assert sum(len(f.text) for f in decoded.fragments) > 2_000_000
    assert decoded.fragments[-1].text.startswith("601")


def test_explicit_unbounded_local_blob_retains_file_over_20mib(tmp_path):
    store = BlobStore(tmp_path / "objects", max_bytes=None)
    payload = b"research\n" * (21 * 1024 * 1024 // 9 + 1)
    result = store.put_stream("long.txt", io.BytesIO(payload))
    assert result.size == len(payload)
    assert result.path.stat().st_size > 20 * 1024 * 1024


def test_profile_reads_all_digest_batches_and_keeps_tail_refs():
    from project_watch.studio.pipeline import build_internal_profile

    class Model:
        def __init__(self):
            self.inputs = []

        def json(self, _system, user, **_kwargs):
            self.inputs.append(user)
            is_tail = "FINAL_EVIDENCE" in user
            return {"summary": "末段资料" if is_tail else "前段资料",
                    "sections": [{"title": "业务与产品", "points": [
                        {"text": "尾部产品" if is_tail else "前部产品",
                         "refs": ["F999" if is_tail else "F1"]}]}]}

    model = Model()
    result = build_internal_profile(model, "Synthetic", "first [F1]\n" * 7000 + "FINAL_EVIDENCE [F999]")
    assert len(model.inputs) > 1
    assert "FINAL_EVIDENCE" in "\n".join(model.inputs)
    assert any(p["refs"] == ["F999"] for s in result["sections"] for p in s["points"])
    assert "末段资料" in result["summary"]


def test_long_fact_keeps_its_reference_in_every_batch():
    from project_watch.studio.pipeline import text_batches

    batches = text_batches("evidence " * 3000 + "FINAL [F999]", 8000)
    assert len(batches) > 1
    assert all("[F999]" in batch for batch in batches)
    assert "FINAL" in batches[-1]


def test_filing_company_match_beyond_page_40_is_found(tmp_path):
    from project_watch.studio.pipeline import filing_passages

    buffer = io.BytesIO()
    canvas = Canvas(buffer)
    for number in range(1, 43):
        canvas.drawString(72, 720, "Synthetic Corp acquired an asset" if number == 42 else "Other evidence")
        canvas.showPage()
    canvas.save()
    assert "[第42页]" in filing_passages(buffer.getvalue(), ["Synthetic Corp"])


def test_internal_model_journal_can_be_unbounded_but_explicit_budget_still_works(tmp_path):
    import pytest
    from test_studio_agent_instance import make_service

    from project_watch.studio.execution import JobExecution
    from project_watch.studio.http import HttpError, Response

    service, pid = make_service(tmp_path / "jobs.sqlite3")
    service.store.enqueue(pid, "internal")
    job = service.store.claim()
    execution = JobExecution(service.store, job)
    for index in range(65):
        execution.call("model", f"synthetic-{index}", lambda: Response(200, {}, b'{}'))
    bounded = JobExecution(service.store, job, budgets={"model": 65})
    with pytest.raises(HttpError) as caught:
        bounded.call("model", "synthetic-next", lambda: pytest.fail("must honor explicit budget"))
    assert caught.value.code == "request_budget"


def test_upload_parse_and_store_preserve_all_large_document_fragments(tmp_path):
    from fastapi.testclient import TestClient

    from project_watch.api import create_app

    app = create_app(tmp_path / "watch.sqlite3", isolated_intake=True)
    client = TestClient(app)
    pid = client.post("/api/projects", json={"name": "Synthetic long document"}).json()["id"]
    payload = "\n".join(f"evidence-{number}: " + "x" * 4000 for number in range(1, 602)).encode()
    uploaded = client.post(f"/api/projects/{pid}/files",
                           files={"file": ("synthetic-long.txt", payload, "text/plain")})
    assert uploaded.status_code == 202
    processed = app.state.intake_worker.process_one()
    assert processed.status == "completed", processed.error_code
    sid = uploaded.json()["source"]["id"]
    fragments = client.get(f"/api/sources/{sid}/fragments").json()
    assert len(fragments) == 601
    assert sum(len(f["text"]) for f in fragments) > 2_000_000
    assert fragments[-1]["text"].startswith("evidence-601:")


def test_changes_process_tail_context_and_persist_all_results(tmp_path, monkeypatch):
    from project_watch.studio import review
    from project_watch.studio.app import create_studio_app

    app = create_studio_app(tmp_path, start_worker=False)
    pipeline = app.state.studio_pipeline
    pid = pipeline.repository.create_project("Synthetic tail changes").id
    metrics = [{"text": f"metric-{n}"} for n in range(45)]
    monkeypatch.setattr(review, "document_changes", lambda *_args: metrics)
    monkeypatch.setattr(pipeline, "_internal_texts", lambda *_args: {"F1": "synthetic source text"})
    monkeypatch.setattr(pipeline, "_internal_cite", lambda marker, ref, text, **_kwargs: {"quote": text, **ref})
    captured = []
    monkeypatch.setattr(pipeline.store, "add_doc_change", lambda *args: captured.append(args[-1]))

    class Model:
        def __init__(self):
            self.inputs = []

        def json(self, _system, user, **_kwargs):
            self.inputs.append(user)
            return {"changes": [{"kind": "新风险", "text": f"risk-{n}", "refs": ["F1"]} for n in range(25)]}

    model = Model()
    refs = {"F1": {"source_id": "new", "source_name": "synthetic.txt", "locator": "line 1"}}
    pipeline._record_changes(pid, model, refs, {"new"},
                             {"summary": "prior\n" * 3000 + "OLD_TAIL"},
                             [{"text": "fresh evidence " * 2000 + "NEW_TAIL", "ref": "F1"}])
    assert "OLD_TAIL" in "\n".join(model.inputs)
    assert "NEW_TAIL" in "\n".join(model.inputs)
    assert captured[0]["metrics"] == metrics
    assert len(captured[0]["texts"]) == 25


def test_evidence_contract_and_exact_reference_resolve_beyond_fragment_500(tmp_path):
    from test_studio_agent_instance import make_service

    from project_watch.studio.evidence_store import EvidenceStore
    from project_watch.studio.monitoring_contracts import EvidenceInput, EvidencePart, LocatedRef

    service, pid = make_service(tmp_path / "evidence.sqlite3")
    parts = [EvidencePart(locator_kind="fragment", locator=f"line {n}", text=f"evidence-{n}:" + "x" * 4000)
             for n in range(1, 602)]
    evidence = EvidenceStore(service.store)
    stored = evidence.save(EvidenceInput(project_id=pid, source_key="synthetic-long",
                                        extraction_version="full-test", read_scope="fragmented", parts=parts))
    ref = LocatedRef(evidence_version_id=stored["id"], part_index=600,
                     start=0, end=13, quote="evidence-601:")
    resolved = evidence.resolve(pid, ref)
    assert resolved["locator"] == "line 601"
    assert len(evidence.get(pid, stored["id"])["parts"]) == 601
