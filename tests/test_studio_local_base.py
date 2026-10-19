"""Synthetic private Studio-only runtime: preserve intake, boundaries and stored data."""
import builtins
from io import BytesIO

import pytest
from fastapi.testclient import TestClient

from project_watch.studio.app import create_studio_app


def test_studio_only_intake_and_history_without_classic_candidates(tmp_path, monkeypatch):
    import project_watch.intake_worker as intake
    def forbidden(*args, **kwargs):
        raise AssertionError('classic candidate analysis must not run')
    for name in ('analyze_document', 'map_financial_fragments', 'map_operating_fragments'):
        monkeypatch.setattr(intake, name, forbidden)
    app = create_studio_app(tmp_path, start_worker=False, classic=False)
    with TestClient(app) as client:
        assert client.get('/').status_code == 200
        assert client.get('/classic').status_code == 503
        assert client.get('/api/projects').status_code == 404
        assert client.get('/api/runtime').json()['classic_available'] is False
        response = client.post('/api/studio/projects', json={
            'entity': {'legal_name': '合成资料公司'}, 'consent': True})
        assert response.status_code == 201, response.text
        pid = response.json()['id']
        response = client.post(f'/api/studio/projects/{pid}/files',
                               files={'file': ('synthetic.txt', b'Synthetic revenue 100', 'text/plain')})
        assert response.status_code == 202, response.text
        worker = app.state.intake_worker
        assert worker.process_one().status == 'completed'
        source, = worker.repository.list_sources(pid)
        fragments = client.get(f'/api/sources/{source.id}/fragments')
        assert fragments.status_code == 200 and fragments.json()[0]['text'] == 'Synthetic revenue 100'
        assert client.get('/api/sources/missing/fragments').status_code == 404
        assert client.get('/api/studio/health').json()['restart_required'] is False
        assert client.patch(f'/api/studio/projects/{pid}', json={'status': 'paused'}).status_code == 200
        assert client.post('/api/studio/runtime/drain').json()['drained']
        settings = app.state.studio_pipeline.settings_store.load()
        assert worker.blob_store.path_for_hash(source.content_sha256).exists()
    reopened = create_studio_app(tmp_path, start_worker=False, classic=False)
    assert reopened.state.studio_store.get_project(pid)['status'] == 'paused'
    assert reopened.state.intake_worker.repository.list_sources(pid)[0].content_sha256 == source.content_sha256
    assert reopened.state.studio_pipeline.settings_store.load() == settings


def test_studio_only_retains_local_origin_protection(tmp_path):
    with TestClient(create_studio_app(tmp_path, start_worker=False, classic=False)) as client:
        response = client.post('/api/studio/projects', headers={'origin': 'https://evil.example'},
                               json={'entity': {'legal_name': '合成公司'}, 'consent': False})
        assert response.status_code == 403
        assert response.headers['x-content-type-options'] == 'nosniff'


def test_broken_classic_dependency_is_not_treated_as_absent(tmp_path, monkeypatch):
    original = builtins.__import__
    def broken(name, *args, **kwargs):
        if name == 'project_watch.api':
            raise ModuleNotFoundError('broken dependency', name='required_dependency')
        return original(name, *args, **kwargs)
    monkeypatch.setattr(builtins, '__import__', broken)
    with pytest.raises(ModuleNotFoundError, match='broken dependency'):
        create_studio_app(tmp_path, start_worker=False)


def test_full_install_keeps_classic_by_default(tmp_path):
    with TestClient(create_studio_app(tmp_path, start_worker=False)) as client:
        assert client.get('/classic').status_code == 200
        assert client.get('/api/projects').status_code == 200


def test_studio_only_background_lifecycle_and_exclusive_home(tmp_path):
    app = create_studio_app(tmp_path, start_worker=True, classic=False)
    with TestClient(app) as client:
        assert app.state.intake_worker.runtime_state() == 'running'
        assert app.state.studio_worker.runtime_state() == 'running'
        assert client.get('/api/studio/health').status_code == 200
        with pytest.raises(RuntimeError, match='已有实例运行'):
            create_studio_app(tmp_path, start_worker=False, classic=False)
    assert app.state.intake_worker.runtime_state() == 'stopped'
    assert app.state.studio_worker.runtime_state() == 'stopped'


def test_full_and_studio_only_share_existing_data_without_migration(tmp_path):
    # Synthetic compatibility only: never open a real user's application home.
    full = create_studio_app(tmp_path, start_worker=False)
    with TestClient(full) as client:
        pid = client.post('/api/studio/projects', json={
            'entity': {'legal_name': 'Synthetic compatibility'}, 'consent': True}).json()['id']
        client.patch(f'/api/studio/projects/{pid}', json={'status': 'paused'})
        assert client.post(f'/api/studio/projects/{pid}/files',
                           files={'file': ('old.txt', b'Existing synthetic original', 'text/plain')}).status_code == 202
        assert full.state.intake_worker.process_one().status == 'completed'
        source, = full.state.intake_worker.repository.list_sources(pid)
    only = create_studio_app(tmp_path, start_worker=False, classic=False)
    with TestClient(only) as client:
        assert only.state.studio_store.get_project(pid)['status'] == 'paused'
        assert client.get(f'/api/sources/{source.id}/fragments').json()[0]['text'] == 'Existing synthetic original'
        assert only.state.intake_worker.blob_store.path_for_hash(source.content_sha256).exists()
    restored = create_studio_app(tmp_path, start_worker=False)
    with TestClient(restored) as client:
        assert client.get('/classic').status_code == 200
        assert restored.state.studio_store.get_project(pid)['status'] == 'paused'
        assert restored.state.intake_worker.repository.get_source(source.id).content_sha256 == source.content_sha256


@pytest.mark.parametrize('suffix', ['pdf', 'docx', 'pptx', 'xlsx'])
def test_studio_only_office_and_pdf_intake(tmp_path, suffix):
    data = BytesIO()
    if suffix == 'pdf':
        from reportlab.pdfgen.canvas import Canvas
        document = Canvas(data)
        document.drawString(72, 700, 'Synthetic source')
        document.save()
    elif suffix == 'docx':
        from docx import Document
        document = Document()
        document.add_paragraph('Synthetic source')
        document.save(data)
    elif suffix == 'pptx':
        from pptx import Presentation
        document = Presentation()
        document.slides.add_slide(document.slide_layouts[1]).shapes.title.text = 'Synthetic source'
        document.save(data)
    else:
        from openpyxl import Workbook
        document = Workbook()
        document.active['A1'] = 'Synthetic source'
        document.save(data)
    app = create_studio_app(tmp_path, start_worker=False, classic=False)
    with TestClient(app) as client:
        pid = client.post('/api/studio/projects', json={
            'entity': {'legal_name': 'Synthetic file company'}, 'consent': True}).json()['id']
        uploaded = client.post(f'/api/studio/projects/{pid}/files',
                               files={'file': (f'synthetic.{suffix}', data.getvalue())})
        assert uploaded.status_code == 202, uploaded.text
        assert app.state.intake_worker.process_one().status == 'completed'
        source, = app.state.intake_worker.repository.list_sources(pid)
        fragments = client.get(f'/api/sources/{source.id}/fragments').json()
        assert any('Synthetic source' in fragment['text'] for fragment in fragments)
