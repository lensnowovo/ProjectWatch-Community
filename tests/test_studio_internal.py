"""Studio internal-document depth: normalisation, integrity checks, old-vs-new documents.

Synthetic material only; the model is a fake transport and network is never used.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from project_watch.studio.http import Http, HttpError, Response
from project_watch.studio.normalize import (
    normalize_amount,
    normalize_metric,
    period_key,
    standard_name,
)
from project_watch.studio.review import document_changes, metric_issues


@pytest.mark.parametrize(("raw", "key"), [
    ("2025", "2025"), ("2025年", "2025"), ("FY2025", "2025"), ("2025E", "2025"), ("2025年(E)", "2025"),
    ("2025年度", "2025"), ("2025H1", "2025H1"), ("2025年上半年", "2025H1"), ("2025年1-6月", "2025H1"),
    ("2025Q3", "2025Q3"), ("2025年第三季度", "2025Q3"), ("3Q2025", "2025Q3"), ("2025年7-9月", "2025Q3"),
    ("2025年1-9月", "2025-1-9M"), ("2025年前三季度", "2025-1-9M"), ("2025-06", "2025M06"),
    ("2025年6月", "2025M06"), ("2025年1-12月", "2025"), ("近三年", "近三年"),
])
def test_period_keys(raw, key):
    assert period_key(raw) == key


@pytest.mark.parametrize(("value", "unit", "context", "expected"), [
    (5800, "万元", "营业收入", (5800.0, "万元", "CNY")),
    (5.8, "亿元", "", (58000.0, "万元", "CNY")),
    (58_000_000, "元", "", (5800.0, "万元", "CNY")),
    (2, "亿", "融资金额", (20000.0, "万元", "CNY")),
    (12, "百万美元", "", (1200.0, "万美元", "USD")),
    (3.2, "million USD", "revenue", (320.0, "万美元", "USD")),
    (35, "%", "", (35, "%", "")),
    (1200, "家", "客户", (1200, "家", "")),
])
def test_amount_normalisation_is_deterministic(value, unit, context, expected):
    assert normalize_amount(value, unit, context) == expected


def test_metric_names_and_ratio_percentages():
    assert standard_name("营收") == standard_name("Revenue") == "营业收入"
    assert standard_name("归母净利润") == "净利润"
    assert standard_name("试用客户数") == "试用客户数"
    margin = normalize_metric({"name": "毛利率", "period": "2025", "value": 0.35, "unit": ""})
    assert margin["value_norm"] == pytest.approx(35) and margin["norm_unit"] == "%"


@pytest.mark.parametrize("unit,context,expected", [
    ("USD 万元", "", (2, "万美元", "USD")),
    ("万美元", "此前人民币收入另列", (2, "万美元", "USD")),
    ("million CNY", "", (200, "万元", "CNY")),
    ("HK$ million", "", (200, "万港元", "HKD")),
])
def test_currency_and_scale_are_kept_together(unit, context, expected):
    assert normalize_amount(2, unit, context) == expected


def test_different_currencies_units_and_profit_scopes_are_not_compared():
    metrics = [_m("old", "营业收入", "2025", 6000, basis="forecast"),
               _m("new", "营业收入", "2025", 1, unit="USD 万元"),
               _m("other", "营业收入", "2025", 2, unit="万港元"),
               _m("old", "净利润", "2025", 1000), _m("new", "归母净利润", "2025", 800)]
    assert not document_changes(metrics, {"new"})
    assert not any("营业收入" in i["text"] and "高于" in i["text"] for i in metric_issues(metrics))
    assert not any("净利润" in i["text"] for i in metric_issues(metrics))
    assert _m("new", "归母净利润", "2025", 800)["scope"] == "归母"


def test_recent_plan_is_the_baseline_even_when_it_equals_the_new_actual():
    oldest = _m("v1", "营业收入", "2025", 6000, basis="forecast")
    latest = _m("v2", "营业收入", "2025", 8000, basis="forecast")
    newest = _m("v3", "营业收入", "2025", 9000)
    oldest["source_created_at"], latest["source_created_at"] = "2025-01-01", "2025-12-01"
    changes = document_changes([oldest, latest, newest], {"v3"})
    assert len(changes) == 1 and "8,000万元" in changes[0]["text"]
    assert [r["source_id"] for r in changes[0]["refs"]] == ["v2", "v3"]
    newest = _m("v3", "营业收入", "2025", 8000)
    assert document_changes([oldest, latest, newest], {"v3"}) == []


def test_quarter_month_range_and_unknown_period_sorting():
    from project_watch.studio.normalize import sort_key

    assert period_key("2025年1-3月") == "2025Q1"
    assert sort_key("2025Q") and sort_key("2025Munknown")


def _m(source, name, period, value, unit="万元", basis="actual", scope="", locator="line 1"):
    return normalize_metric({"source_id": source, "source_name": f"{source}.pdf", "name": name,
                             "period": period, "value": value, "unit": unit, "basis": basis,
                             "scope": scope, "locator": locator})


def test_integrity_issues():
    metrics = [
        _m("a", "营业收入", "2025", 5800), _m("b", "营收", "2025年", 5.2, unit="亿元"),  # 5800 vs 52000
        _m("a", "净利润", "2024", 4000), _m("a", "营业收入", "2024", 3200),
        _m("a", "毛利率", "2025", 135, unit="%"),
        _m("a", "营业收入", "2023", 1000), _m("b", "营业收入", "2023", 0.1, unit="亿元"),  # same number
    ]
    issues = metric_issues(metrics)
    texts = " ".join(i["text"] for i in issues)
    assert "营业收入 2025" in texts and "不一致" in texts
    assert "2024 净利润" in texts and "高于营业收入" in texts
    assert "毛利率 2025" in texts
    assert "2023" not in texts  # 1000万元 == 0.1亿元, not a conflict
    conflict = next(i for i in issues if i["kind"] == "conflict")
    assert {r["source_id"] for r in conflict["refs"]} == {"a", "b"}


def test_document_changes_compare_new_against_earlier():
    old = [_m("v1", "营业收入", "2025", 6000, basis="forecast"), _m("v1", "营业收入", "2026", 12000, basis="forecast"),
           _m("v1", "净利润", "2024", -1500), _m("v1", "货币资金", "2024", 3000)]
    new = [_m("v2", "营业收入", "2025年", 5800, scope="合并"), _m("v2", "营业收入", "2026E", 15000, basis="forecast"),
           _m("v2", "净利润", "2024", -1600), _m("v2", "货币资金", "2025", 4200),
           _m("v2", "营业收入", "2027", 20000, basis="forecast")]
    changes = document_changes(old + new, {"v2"})
    kinds = [c["kind"] for c in changes]
    assert kinds[:1] == ["actual_vs_plan"]
    vs_plan = changes[0]
    assert "实际 5,800万元" in vs_plan["text"] and "预测 6,000万元" in vs_plan["text"] and "-3%" in vs_plan["text"]
    assert [r["source_id"] for r in vs_plan["refs"]] == ["v1", "v2"]
    assert any(c["kind"] == "plan_revised" and "+25%" in c["text"] for c in changes)
    assert any(c["kind"] == "restated" and "净利润 2024" in c["text"] for c in changes)
    assert any(c["kind"] == "new_period" and "货币资金 新增 2025" in c["text"] for c in changes)
    assert any(c["kind"] == "new_period" and "2027" in c["text"] for c in changes)
    assert document_changes(new, {"v2"}) == []  # nothing to compare against


def test_pptx_tables_keep_headers(tmp_path: Path):
    from pptx import Presentation
    from pptx.util import Inches

    from project_watch.extractors import extract_document

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(3, 3, Inches(1), Inches(1), Inches(6), Inches(2)).table
    for r, row in enumerate([["指标（万元）", "2024", "2025"], ["营业收入", "3200", "5800"], ["净利润", "-1500", "-900"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    path = tmp_path / "bp.pptx"
    deck.save(path)
    text = extract_document(path).fragments[0].text
    assert "营业收入 | 2024: 3200 | 2025: 5800" in text
    assert "表头: 指标（万元） | 2024 | 2025" in text


# ---------------------------------------------------------------- end to end

BP_V1 = """光启视界 商业计划书 2025年3月版（合成数据）
公司专注车载激光雷达光学模组。
营业收入：2024年 3200（实际）；2025年 6000（预测）
客户：两家整车厂试用中。
"""
BP_V2 = """光启视界 经营更新 2026年1月（合成数据）
营业收入：2025年 5800（实际）；2026年 15000（预测）
客户：其中一家整车厂已完成验收并下达量产订单。
新风险：上游芯片交期拉长至 40 周。
"""


class FakeModel:
    def __init__(self) -> None:
        self.prompts: list[str] = []
        self.profile_inputs: list[str] = []

    def answer(self, system: str, user: str):
        self.prompts.append(system)
        if "正在阅读" in system:
            metrics = []
            for marker, line in re.findall(r"\[(F\d+)\][^\n]*?\) ([^\n]*)", user):
                name = line.split("：")[0] if "：" in line else ""
                for year, value, basis in re.findall(r"(\d{4})年 ([\d.,-]+)（(实际|预测)）", line):
                    metrics.append({"name": name, "category": "financial", "period": year, "value": value,
                                    "unit": "万元", "basis": "actual" if basis == "实际" else "forecast",
                                    "ref": marker, "quote": line[:40]})
            first = re.findall(r"\[(F\d+)\]", user)
            return {"points": [{"section": "经营进展", "text": line, "ref": m}
                               for m, line in re.findall(r"\[(F\d+)\][^\n]*?\) ([^\n]*)", user)],
                    "metrics": metrics, "milestones": [], "risks": [], "financing": [], "_": first}
        if "一页纸档案" in system:
            self.profile_inputs.append(user)
            return {"one_liner": "合成", "summary": "公司专注光学模组。公司预计2025年营业收入6000万元。",
                    "sections": [{"title": "经营进展", "points": [
                        {"text": "2025年收入目标 6000 万元", "refs": ["F1"]},
                        {"text": "两家整车厂试用中", "refs": ["F1"]}]}],
                    "highlights": [], "risks": [],
                    "milestones": [{"date": "2025", "text": "实现收入6000万元", "status": "planned", "refs": ["F1"]}],
                    "questions": []}
        if "新一批内部资料" in system:
            refs = re.findall(r"\[(F\d+)\]", user.split("## 新资料中的要点")[1])
            return {"changes": [
                {"kind": "新进展", "text": "一家整车厂已完成验收并下达量产订单", "refs": refs[:1]},
                {"kind": "新风险", "text": "无出处的说法应被丢弃", "refs": ["F999"]}]}
        if "动态简报" in system:
            return {"headline": "预计2025年收入6000万元", "tone": "neutral",
                    "summary": "预计2025年收入6000万元。客户验证在推进。",
                    "changes": [{"text": "实际2025年收入5800万元，低于此前预测6000万元。", "refs": []}],
                    "watch": [{"text": "预计2025年收入6000万元。关注客户验收。", "refs": []}]}
        if "最值得持续跟踪的" in system:
            return {"questions": []}
        if "下面每个问题附带了本轮新命中的事件" in system:
            return {"results": []}
        raise AssertionError(system[:30])

    def __call__(self, method, url, headers, body, timeout):
        if url.startswith("https://model.test/v1/chat/completions"):
            payload = json.loads(body)
            data = self.answer(payload["messages"][0]["content"], payload["messages"][1]["content"])
            return Response(200, {}, json.dumps({"choices": [{"message": {"content": json.dumps(data, ensure_ascii=False)}}],
                                                 "usage": {"prompt_tokens": 1, "completion_tokens": 1}}).encode())
        raise HttpError("connection_failed")


def _drain(app):
    for _ in range(40):
        progressed = app.state.intake_worker.process_one() is not None
        worker = app.state.studio_worker
        worker.schedule()
        progressed = worker.step() or progressed
        if not progressed:
            return


def test_second_document_produces_change_set(tmp_path):
    from project_watch.studio.app import create_studio_app

    model = FakeModel()
    app = create_studio_app(tmp_path, start_worker=False)
    app.state.studio_pipeline.http_factory = lambda settings: Http(model)
    client = TestClient(app)
    client.put("/api/studio/settings", json={"model": {
        "protocol": "openai", "base_url": "https://model.test/v1", "model": "m", "api_key": "k-123456789"}})
    project = client.post("/api/studio/projects", json={"entity": {"legal_name": "光启视界"}, "consent": True}).json()
    store = app.state.studio_store
    # Only internal analysis matters here: drop the queued external job.
    with store._connect() as db:
        # Remove only admission metadata in this synthetic queue-reset fixture.
        db.execute("DELETE FROM studio_job_checkpoints WHERE stage='collection_request_v1'")
        db.execute("DELETE FROM studio_jobs")

    client.post(f"/api/studio/projects/{project['id']}/files", files={"file": ("bp-v1.txt", BP_V1.encode(), "text/plain")})
    _drain(app)
    detail = client.get(f"/api/studio/projects/{project['id']}").json()
    assert detail["doc_changes"] == []
    assert "6000万元" in detail["profile"]["summary"]  # still the latest view after v1
    revenue = {m["period_key"]: m for m in detail["metrics"] if m["std_name"] == "营业收入"}
    assert revenue["2025"]["basis"] == "forecast" and revenue["2025"]["value_norm"] == 6000

    client.post(f"/api/studio/projects/{project['id']}/files", files={"file": ("update-v2.txt", BP_V2.encode(), "text/plain")})
    _drain(app)
    detail = client.get(f"/api/studio/projects/{project['id']}").json()
    [change] = detail["doc_changes"]
    assert change["documents"] == ["update-v2.txt"]
    texts = [c["text"] for c in change["metrics"]]
    assert any("营业收入 2025 实际 5,800万元，此前预测 6,000万元" in t for t in texts)
    assert any("2026" in t and "新增" in t for t in texts)
    # The fixture cites its first (financial) fragment, not the actual order
    # fragment. Keep the proposal inspectable but do not borrow a financial
    # quote or claim that this is a source-supported new event.
    assert change['texts'][0]['kind'] == '资料补充'
    assert change['texts'][0]['text'] == '新增资料，未定位到对应的新进展原句。'
    assert '下达量产订单' in change['texts'][0]['proposed_text']
    assert change['texts'][0]['evidence_support'] == 'source_only'
    assert not change['texts'][0]['refs'][0].get('quote')
    assert change["texts"][0]["refs"][0]["source_name"] == "update-v2.txt"
    assert detail["issues"] == []
    # The profile no longer presents the replaced 2025 forecast as current.
    assert "已被实际数取代" in model.profile_inputs[-1] and "资料2" in model.profile_inputs[-1]
    profile = detail["profile"]
    assert profile["summary"] == "公司专注光学模组。" and profile["stale_removed"] == 3
    assert [p["text"] for s in profile["sections"] for p in s["points"]] == ["两家整车厂试用中"]
    assert profile["milestones"] == []
    assert detail["brief"]["headline"] == ""
    assert detail["brief"]["summary"] == "客户验证在推进。"
    assert detail["brief"]["watch"][0]["text"] == "关注客户验收。"
    assert "此前预测6000万元" in detail["brief"]["changes"][0]["text"]
    assert detail["brief"]["stale_removed"] == 3


def test_metrics_written_before_normalisation_are_backfilled_on_read(tmp_path):
    from project_watch.studio.app import create_studio_app

    store = create_studio_app(tmp_path, start_worker=False).state.studio_store
    with store._connect() as db:
        db.execute("PRAGMA foreign_keys=OFF")  # the projects table belongs to the main schema
        db.execute("""INSERT INTO studio_metrics (id, project_id, source_id, name, category, period, value, unit,
                      basis, scope, locator, quote, created_at) VALUES
                      ('m','p','s','营收','financial','2025年',0.58,'亿元','actual','','page 1','', 't')""")
    [metric] = store.list_metrics("p")
    assert (metric["std_name"], metric["period_key"], metric["value_norm"], metric["norm_unit"]) == ("营业收入", "2025", 5800, "万元")


def test_stale_forecast_sentences_are_scrubbed_but_comparisons_survive():
    from project_watch.studio.review import current_lines, scrub_stale, superseded

    metrics = [_m("a", "营业收入", "2025E", 6500, basis="forecast"), _m("b", "营收", "2025年", 0.58, unit="亿元"),
               _m("b", "营业收入", "2026E", 15000, basis="forecast")]
    replaced = superseded(metrics)
    assert current_lines(metrics) == ["- 营业收入 2025 5,800万元（实际）；此前预测 6,500万元，已被实际数取代",
                                      "- 营业收入 2026 15,000万元（预测）"]
    text = "公司预计2025年收入6500万元，同比翻倍。2025年实际收入5800万元，低于此前预测的6500万元。2026年目标1.5亿元。"
    assert scrub_stale(text, replaced) == ("2025年实际收入5800万元，低于此前预测的6500万元。2026年目标1.5亿元。", 1)
    assert scrub_stale("2025年营收预计0.65亿元。", replaced) == ("", 1)
    assert scrub_stale("公司预计2025年交付 6500 台。", replaced)[1] == 0  # another metric/unit is not a stale revenue
    assert scrub_stale("公司预计2027年收入6500万元。", replaced)[1] == 0  # different year
    assert scrub_stale("2025年收入6500万元。", []) == ("2025年收入6500万元。", 0)


@pytest.mark.parametrize("sentence", [
    "预计2025年研发投入6500万元。", "预计2025年收入6500万美元。",
    "预计2025年上半年收入6500万元。", "预计2025年母公司收入6500万元。",
    "预计2025年毛利率65%。",
])
def test_stale_scrub_preserves_other_metrics_currencies_periods_and_scopes(sentence):
    from project_watch.studio.review import scrub_stale, superseded

    replaced = superseded([_m("a", "营业收入", "2025", 6500, basis="forecast"),
                           _m("b", "营业收入", "2025", 5800)])
    assert scrub_stale(sentence, replaced) == (sentence, 0)


def test_current_lines_uses_latest_old_plan():
    from project_watch.studio.review import current_lines

    metrics = [_m("a", "营业收入", "2025", 6000, basis="forecast") | {"source_created_at": "2025-01-01"},
               _m("b", "营业收入", "2025", 6500, basis="forecast") | {"source_created_at": "2025-06-01"},
               _m("c", "营业收入", "2025", 5800) | {"source_created_at": "2026-01-01"}]
    [line] = current_lines(metrics)
    assert "6,500万元" in line and "6,000万元" not in line


def test_profile_scrub_keeps_valid_sentences_in_the_same_point():
    from project_watch.studio.pipeline import Pipeline
    from project_watch.studio.review import superseded

    replaced = superseded([_m("a", "营业收入", "2025", 6000, basis="forecast"),
                           _m("b", "营业收入", "2025", 5800)])
    result = {"summary": "实际5800万元，低于此前预测6000万元。", "one_liner": "合成光学公司",
              "sections": [{"points": [{"text": "预计2025年收入6000万元。客户验收完成。"}]}],
              "highlights": [], "milestones": [], "risks": [{"text": "预计2025年收入6000万元。芯片交期较长。"}],
              "questions": ["预计2025年收入6000万元？", "2026年目标如何实现？"]}
    assert Pipeline._scrub_profile(result, replaced) == 3
    assert result["sections"][0]["points"][0]["text"] == "客户验收完成。"
    assert result["risks"][0]["text"] == "芯片交期较长。"
    assert result["questions"] == ["2026年目标如何实现？"]


# ---------------------------------------------------------------- external: reading fallbacks

class _Reader:
    def __init__(self) -> None:
        self.urls: list[str] = []

    def read(self, url):
        self.urls.append(url)
        if url == "https://direct.example/a":
            return {"status": "article_body", "text": "光启视界宣布完成3亿元B轮融资，红杉中国领投，资金用于苏州二期产线。"}
        return {"status": "summary_only", "text": "", "error_code": "access_restricted"}


def test_unreadable_reports_fall_back_to_twin_text_or_search_summary(tmp_path):
    from project_watch.repository import ProjectRepository
    from project_watch.studio.pipeline import Pipeline
    from project_watch.studio.settings import SettingsStore
    from project_watch.studio.store import StudioStore

    repo = ProjectRepository(tmp_path / "db.sqlite3")
    project = repo.create_project("光启视界")
    store = StudioStore(repo.db_path)
    reader = _Reader()
    pipeline = Pipeline(repo, store, SettingsStore(tmp_path / "settings.json"), article_reader=reader)
    long_summary = "光启视界与某新势力车企签署量产定点协议，" * 8
    store.add_items(project.id, "company", [
        {"key": "g1", "title": "光启视界完成3亿元B轮融资 红杉中国领投", "url": "https://news.google.com/rss/articles/1",
         "source": "36氪", "snippet": "", "published_at": "2026-10-06T00:00:00Z"},
        {"key": "g2", "title": "光启视界创始人谈降本", "url": "https://news.google.com/rss/articles/2",
         "source": "钛媒体", "snippet": "", "published_at": "2026-10-06T00:00:00Z"},
        {"key": "d1", "title": "光启视界完成3亿元B轮融资，红杉中国领投", "url": "https://direct.example/a",
         "source": "direct.example", "snippet": "", "published_at": "2026-10-06T00:00:00Z"},
        {"key": "b1", "title": "光启视界获定点", "url": "https://walled.example/b", "source": "博查",
         "snippet": long_summary, "published_at": "2026-10-05T00:00:00Z"},
        {"key": "b2", "title": "光启视界短讯", "url": "https://walled.example/c", "source": "博查",
         "snippet": "短摘要", "published_at": "2026-10-05T00:00:00Z"},
        {"key": "old", "title": "光启视界完成3亿元B轮融资 红杉中国领投", "url": "https://news.google.com/rss/articles/old",
         "source": "36氪", "snippet": "", "published_at": "2026-08-01T00:00:00Z"},
        {"key": "conflict", "title": "光启视界完成5亿元B轮融资 红杉中国领投", "url": "https://news.google.com/rss/articles/conflict",
         "source": "36氪", "snippet": "", "published_at": "2026-10-06T00:00:00Z"},
    ])
    assert pipeline.fetch_bodies(project.id, lambda _: None) == 1
    assert "news.google.com" not in " ".join(reader.urls)
    by_key = {i["key"]: i for i in store.untriaged(project.id)}
    assert by_key["g1"]["body_status"] == "twin_body" and "红杉中国领投" in by_key["g1"]["body"]
    assert json.loads(by_key["g1"]["extra"])["twin_source"]["id"] == by_key["d1"]["id"]
    assert by_key["g2"]["body_status"] == "skipped"
    assert by_key["b1"]["body_status"] == "search_summary" and by_key["b1"]["body"] == long_summary.strip()[:1200]
    assert by_key["b2"]["body_status"] == "summary_only" and by_key["b2"]["body"] == ""
    assert by_key["old"]["body_status"] == by_key["conflict"]["body_status"] == "skipped"
    store.add_items(project.id, "company", [{"key": "pending", "title": "待阅读报道", "url": "https://pending.example/a",
        "source": "pending", "snippet": "", "published_at": None}])
    store.add_items(project.id, "paper", [{"key": "paper", "title": "论文", "url": "https://paper.example/a",
        "source": "OpenAlex", "snippet": "", "published_at": None}])

    coverage = store.body_coverage(project.id)
    assert {k: coverage[k] for k in ("total", "article", "twin", "summary", "none")} == {
        "total": 8, "article": 1, "twin": 1, "summary": 1, "none": 5}
    assert coverage["outlets"] == [{"host": "walled.example", "read": 0, "total": 2}]


def test_body_budget_does_not_hide_unread_reports_from_coverage(tmp_path, monkeypatch):
    from project_watch.repository import ProjectRepository
    from project_watch.studio import pipeline as module
    from project_watch.studio.settings import SettingsStore
    from project_watch.studio.store import StudioStore

    repo = ProjectRepository(tmp_path / "db.sqlite3")
    project = repo.create_project("光启视界")
    store = StudioStore(repo.db_path)
    long_summary = "公开搜索摘要，只用于测试兜底标签。" * 10
    store.add_items(project.id, "company", [
        {"key": str(i), "title": f"新闻{i}", "url": f"https://budget.example/{i}", "source": "搜索",
         "snippet": snippet, "published_at": None} for i, snippet in enumerate([long_summary, "短摘要"])])
    monkeypatch.setattr(module, "BODY_FETCH_LIMIT", 0)
    pipeline = module.Pipeline(repo, store, SettingsStore(tmp_path / "settings.json"), article_reader=_Reader())
    assert pipeline.fetch_bodies(project.id, lambda _: None) == 0
    assert {r["body_status"] for r in store.untriaged(project.id)} == {"search_summary", "not_read"}
    coverage = store.body_coverage(project.id)
    assert coverage["total"] == 2 and coverage["summary"] == coverage["none"] == 1


def test_event_link_prefers_read_article_over_google_redirect(tmp_path):
    from project_watch.repository import ProjectRepository
    from project_watch.studio.store import StudioStore

    repo = ProjectRepository(tmp_path / "db.sqlite3")
    project = repo.create_project("光启视界")
    store = StudioStore(repo.db_path)
    store.add_items(project.id, "company", [
        {"key": "g", "title": "光启视界完成B轮融资", "url": "https://news.google.com/rss/articles/1", "source": "36氪",
         "snippet": "", "published_at": "2026-10-05T00:00:00Z"},
        {"key": "d", "title": "光启视界完成B轮融资", "url": "https://direct.example/a", "source": "direct.example",
         "snippet": "", "published_at": "2026-10-06T00:00:00Z"}])
    items = {i["key"]: i for i in store.untriaged(project.id)}
    store.save_body(items["d"]["id"], "article_body", "正文")
    store.save_triage([{"id": i["id"], "relevant": True, "importance": "major", "summary": "融资"} for i in items.values()])
    store.create_event(project.id, {"title": "光启视界完成B轮融资", "category": "company", "importance": "major",
                                    "summary": "融资", "impact": "", "event_at": "2026-10-05T00:00:00Z"},
                       [i["id"] for i in items.values()])
    [event] = store.list_events(project_id=project.id)
    assert event["url"] == "https://direct.example/a" and event["source_count"] == 2


@pytest.mark.parametrize("reply", [[], {"changes": "bad"}, {"changes": []}])
def test_empty_changes_publish_but_malformed_changes_preserve_previous_result(tmp_path, reply):
    from project_watch.studio.app import create_studio_app
    from project_watch.studio.llm import ModelError

    app = create_studio_app(tmp_path, start_worker=False)
    repository = app.state.intake_worker.repository
    project = repository.create_project("合成变化项目")
    pipeline = app.state.studio_pipeline
    app.state.studio_store.add_doc_change(project.id, ["old"],
                                         {"documents": ["old.txt"], "metrics": [], "texts": [{"text": "旧变化"}]})

    class Model:
        def json(self, *_args, **_kwargs):
            return reply

    def compare():
        pipeline._record_changes(project.id, Model(),
            {"F1": {"source_id": "new", "source_name": "new.txt", "locator": "line 1"}}, {"new"},
            {"summary": "合成旧资料"}, [{"ref": "F1", "text": "合成新资料"}])

    if reply != {"changes": []}:
        with pytest.raises(ModelError) as failure:
            compare()
        assert failure.value.code == "model_output_invalid_shape"
        [previous] = app.state.studio_store.doc_changes(project.id)
        assert previous["documents"] == ["old.txt"]
        assert previous["texts"] == [{"text": "旧变化"}]
        return
    compare()
    latest = app.state.studio_store.doc_changes(project.id)[0]
    assert latest["documents"] == ["new.txt"]
    assert latest["metrics"] == latest["texts"] == []


def test_items_the_prefilter_declined_are_counted_apart_from_unreadable(tmp_path):
    import json as _json
    from contextlib import closing

    from project_watch.repository import ProjectRepository
    from project_watch.studio.store import StudioStore

    repo = ProjectRepository(tmp_path / "db.sqlite3")
    project = repo.create_project("光启视界")
    store = StudioStore(repo.db_path)
    store.add_items(project.id, "company", [{"key": k, "title": k, "url": f"https://x.example/{k}", "source": "s",
                                             "snippet": "", "published_at": None} for k in ("same-name", "pending")])
    with closing(store._connect()) as db, db:
        db.execute("UPDATE studio_items SET extra=? WHERE key='same-name'",
                   (_json.dumps({"body_prefilter": {"version": "relevance-v1", "read": False}}),))
    coverage = store.body_coverage(project.id)
    assert coverage["total"] == 2 and coverage["declined"] == 1 and coverage["none"] == 1
