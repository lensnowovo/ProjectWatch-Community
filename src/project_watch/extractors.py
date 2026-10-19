"""Complete text extraction in located fragments, never business conclusions."""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import time
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

import pdfplumber
import xlrd
from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pypdf import PdfReader

from project_watch.domain import FragmentInput
from project_watch.file_validation import FileValidationError, validate_document

# Fragment size is a processing unit, not a whole-document completion ceiling.
MAX_FRAGMENT_CHARS = 20_000
OCR_PAGE_TIMEOUT_SECONDS = 8
_COLUMN_HEADER_HINT = re.compile(
    r"(?:19|20)\d{2}|实际|预测|预算|规划|historical|actual|forecast|budget|plan",
    re.IGNORECASE,
)
_NUMERIC_TEXT = re.compile(
    r"(?:[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?"
    r"|\(\s*[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?\s*\))\Z"
)
_SIMPLE_PERCENT_FORMAT = re.compile(r"(?:#,##)?0(?:\.[0#]+)?%\Z")


class ExtractionError(ValueError):
    def __init__(self, error_code: str) -> None:
        self.error_code = error_code
        super().__init__(error_code)


@dataclass(frozen=True, slots=True)
class ExtractionResult:
    fragments: list[FragmentInput]
    needs_ocr: bool


class _Collector:
    def __init__(self, progress=None) -> None:
        self.progress = progress
        self.fragments: list[FragmentInput] = []
        self.total_chars = 0
        self.needs_ocr = False

    def add(self, locator: str, text: str | None) -> None:
        if self.progress:
            self.progress()
        normalized = (text or "").strip()
        if not normalized:
            return
        self.total_chars += len(normalized)
        chunks = [
            normalized[index : index + MAX_FRAGMENT_CHARS]
            for index in range(0, len(normalized), MAX_FRAGMENT_CHARS)
        ]
        for index, chunk in enumerate(chunks, start=1):
            part_locator = f"{locator} part {index}" if len(chunks) > 1 else locator
            self.fragments.append(
                FragmentInput(ordinal=len(self.fragments) + 1, locator=part_locator, text=chunk)
            )

    def finish(self) -> ExtractionResult:
        if not self.fragments and not self.needs_ocr:
            raise ExtractionError("no_text")
        return ExtractionResult(fragments=self.fragments, needs_ocr=self.needs_ocr)


def _extract_text(path: Path, collector: _Collector) -> None:
    payload = path.read_bytes()
    if b"\x00" in payload[:4096]:
        raise ExtractionError("invalid_document")
    try:
        decoded = payload.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            decoded = payload.decode("gb18030")
        except UnicodeDecodeError as exc:
            raise ExtractionError("invalid_document") from exc
    for index, line in enumerate(decoded.splitlines(), start=1):
        collector.add(f"line {index}", line)


def _extract_pdf(path: Path, collector: _Collector) -> None:
    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted:
        raise ExtractionError("encrypted_document")
    # PDFMiner's font mapping handles Chinese subset fonts that can otherwise
    # decode to unrelated Unicode scripts while still appearing "non-empty".
    with pdfplumber.open(path) as document:
        for index, page in enumerate(reader.pages, start=1):
            native_page = document.pages[index - 1]
            try:
                text = native_page.extract_text() or ""
            finally:
                native_page.close()
            # A scan can have only a page number as a text layer. Try OCR on such
            # pages too, preserving the original text if OCR cannot improve it.
            if text.strip() and (len(text.strip()) >= 40 or not page.images):
                collector.add(f"page {index}", text)
            else:
                # Every scan page gets an attempt, including a long report's tail.
                ocr_text = _ocr_pdf_page(path, index, 2 * OCR_PAGE_TIMEOUT_SECONDS)
                if ocr_text and len(ocr_text.strip()) > len(text.strip()):
                    collector.add(f"page {index} OCR", ocr_text)
                else:
                    collector.add(f"page {index}", text)
                    collector.needs_ocr = True


def _ocr_pdf_page(path: Path, page_number: int, budget_seconds: float) -> str | None:
    """Best-effort local OCR; all processes and temporary data stay in parser isolation."""
    from project_watch.local_ocr import recognize_pdf_page

    try:
        text = recognize_pdf_page(path, page_number)
        if text:
            return text
    except Exception:
        # Preserve native text and explicitly flag the gap if both engines fail.
        pass
    if not shutil.which("pdftoppm") or not shutil.which("tesseract"):
        return None
    try:
        started = time.monotonic()
        with tempfile.TemporaryDirectory(prefix="project-watch-ocr-") as directory:
            prefix = Path(directory) / "page"
            rendered = subprocess.run(
                [
                    "pdftoppm",
                    "-f",
                    str(page_number),
                    "-l",
                    str(page_number),
                    "-scale-to",
                    "2200",
                    "-gray",
                    "-singlefile",
                    "-png",
                    str(path),
                    str(prefix),
                ],
                stdin=subprocess.DEVNULL,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=min(OCR_PAGE_TIMEOUT_SECONDS, budget_seconds),
                check=False,
            )
            image = prefix.with_suffix(".png")
            if (
                rendered.returncode
                or not image.is_file()
                or image.stat().st_size > 16 * 1024 * 1024
            ):
                return None
            remaining = budget_seconds - (time.monotonic() - started)
            if remaining <= 0:
                return None
            recognized = subprocess.run(
                ["tesseract", str(image), "stdout", "-l", "chi_sim+eng", "--psm", "11"],
                stdin=subprocess.DEVNULL,
                stdout=subprocess.PIPE,
                stderr=subprocess.DEVNULL,
                timeout=min(OCR_PAGE_TIMEOUT_SECONDS, remaining),
                check=False,
            )
            if recognized.returncode or len(recognized.stdout) > 100_000:
                return None
            return recognized.stdout.decode("utf-8", errors="replace").strip() or None
    except (OSError, subprocess.TimeoutExpired):
        return None


def _extract_docx(path: Path, collector: _Collector) -> None:
    document = Document(str(path))
    for index, paragraph in enumerate(document.paragraphs, start=1):
        collector.add(f"paragraph {index}", paragraph.text)
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            collector.add(f"table {table_index} row {row_index}", text)
    if not collector.fragments:
        collector.needs_ocr = True


def _extract_pptx(path: Path, collector: _Collector) -> None:
    presentation = Presentation(str(path))
    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                parts.append(shape.text.strip())
            if shape.has_table:
                parts.extend(_pptx_table_lines(shape.table))
        if parts:
            collector.add(f"slide {slide_index}", "\n".join(parts))
        else:
            collector.needs_ocr = True


def _pptx_table_lines(table) -> list[str]:
    """Keep the column header next to each value so a number never loses its period/label."""
    rows = [[" ".join(cell.text.split()) for cell in row.cells] for row in table.rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    header = rows[0]
    # Two-column (label | value) tables are unambiguous as plain rows; only tables
    # with several value columns (periods, scenarios) need the header beside each value.
    if len(rows) < 2 or sum(1 for h in header[1:] if h) < 2:
        return [" | ".join(v for v in r if v) for r in rows]
    lines = ["表头: " + " | ".join(h for h in header if h)]
    for row in rows[1:]:
        label = row[0] if row else ""
        pairs = [f"{header[i] or f'列{i + 1}'}: {value}" for i, value in enumerate(row[1:], start=1)
                 if value and i < len(header)]
        extra = [value for value in row[len(header):] if value]
        lines.append(" | ".join(filter(None, [label, *pairs, *extra])))
    return lines


def _extract_xlsx(path: Path, collector: _Collector) -> None:
    # Content-addressed originals have no filename suffix. Loading a binary
    # stream avoids openpyxl's path-extension check without changing storage.
    with path.open("rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in workbook.worksheets:
                headers: dict[int, str] = {}
                for row_index, row in enumerate(sheet.iter_rows(values_only=False), start=1):
                    values = [cell.value for cell in row]
                    percent_columns = {
                        column
                        for column, cell in enumerate(row, start=1)
                        if isinstance(cell.value, (int, float, Decimal))
                        and not isinstance(cell.value, bool)
                        and _SIMPLE_PERCENT_FORMAT.fullmatch(cell.number_format.strip())
                    }
                    cells = _located_sheet_cells(
                        values, row_index, headers, percent_columns=percent_columns
                    )
                    collector.add(f"sheet {sheet.title} row {row_index}", " | ".join(cells))
        finally:
            workbook.close()


def _extract_xls(path: Path, collector: _Collector) -> None:
    workbook = xlrd.open_workbook(str(path), on_demand=True)
    try:
        for sheet in workbook.sheets():
            headers: dict[int, str] = {}
            for row_index in range(sheet.nrows):
                values = []
                for column in range(sheet.ncols):
                    cell = sheet.cell(row_index, column)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        values.append(None)
                        continue
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(cell.value, workbook.datemode).isoformat()
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = "TRUE" if cell.value else "FALSE"
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(cell.value, "#ERROR")
                    else:
                        value = cell.value
                    values.append(value)
                cells = _located_sheet_cells(values, row_index + 1, headers)
                collector.add(f"sheet {sheet.name} row {row_index + 1}", " | ".join(cells))
    finally:
        workbook.release_resources()


def _located_sheet_cells(
    values: tuple[object, ...] | list[object],
    row_index: int,
    headers: dict[int, str],
    *,
    percent_columns: set[int] | None = None,
) -> list[str]:
    """Keep column labels next to values; never infer actual/forecast as a fact."""
    cells: list[str] = []
    updates: dict[int, str] = {}
    for column, value in enumerate(values, start=1):
        if value is None or not str(value).strip():
            continue
        address = f"{get_column_letter(column)}{row_index}"
        hint = headers.get(column)
        numeric_value = (
            isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)
        ) or (
            isinstance(value, str)
            and len(value) <= 60
            and _NUMERIC_TEXT.fullmatch(value.strip()) is not None
        )
        header = f"[header:{hint}]" if hint and numeric_value else ""
        percent_format = (
            "[format:percent]"
            if percent_columns
            and column in percent_columns
            and isinstance(value, (int, float, Decimal))
            and not isinstance(value, bool)
            else ""
        )
        cells.append(f"{address}{header}{percent_format}={value}")
        if isinstance(value, str) and not numeric_value and len(value) <= 60 and _COLUMN_HEADER_HINT.search(value):
            updates[column] = value.strip()
    headers.update(updates)
    return cells


def extract_document(
    path: Path | str, *, suffix: str | None = None, max_bytes: int | None = None, progress=None
) -> ExtractionResult:
    """Extract all located text; preserve uncertainty when OCR is needed."""

    path = Path(path)
    suffix = suffix.lower() if suffix is not None else path.suffix.lower()
    try:
        validate_document(path, suffix, max_bytes=max_bytes)
    except FileValidationError as exc:
        raise ExtractionError(exc.error_code) from exc

    collector = _Collector(progress)
    try:
        if suffix == ".txt":
            _extract_text(path, collector)
        elif suffix == ".pdf":
            _extract_pdf(path, collector)
        elif suffix == ".docx":
            _extract_docx(path, collector)
        elif suffix == ".pptx":
            _extract_pptx(path, collector)
        elif suffix == ".xlsx":
            _extract_xlsx(path, collector)
        elif suffix == ".xls":
            _extract_xls(path, collector)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("invalid_document") from exc
    return collector.finish()
