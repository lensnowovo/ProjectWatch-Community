"""Installed-package combined journey. Built-in synthetic inputs, no network.

Reuses the existing intake/consent/reading smoke, then exercises external focus
and source association in that SAME project. It cannot certify model quality or
public coverage and never opens the pilot DB. --serve is an isolated UI fixture.
"""
import argparse
import io
import json
import os
import runpy
import tempfile
from datetime import UTC, datetime, timedelta
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from project_watch.api import create_app
from project_watch.monitoring import FeedItem, ManualSnapshotAdapter
from project_watch.public_search import gdelt_rss_target
from project_watch.repository import ProjectRepository


def exercise(root):
    fixture = runpy.run_path(str(Path(__file__).with_name('smoke_company_reading.py')))
    report = dict(passed=False, private_materials_used=False, real_model_quality='not_tested',
                  real_public_collection='not_tested',
                  independent_use_gate='not_passed_operator_seeded', checks={})

    def check(name, value):
        report['checks'][name] = bool(value)
        assert value, name

    with patch.dict(os.environ, {'PROJECT_WATCH_SETTINGS_TOKEN':'synthetic-local-only'}):
        baseline_app, project_id = fixture['exercise'](root)
        baseline_app.state.reading_worker.stop()
        model = baseline_app.state.reading_worker.settings.get()
        db = Path(root) / 'watch.sqlite3'
        repo = ProjectRepository(db)
        now = datetime.now(UTC)
        term = 'Synthetic Peer'
        feed = (
            FeedItem('1', f'{term} 公布季度营收', 'https://example.org/journey/results',
                     '合成新闻样本，不是实际经营信息。', now),
            FeedItem('2', f'{term}less 公布季度营收', 'https://example.org/journey/wrong', '', now),
            FeedItem('3', f'{term} 公布历史营收', 'https://example.org/journey/old', '', now-timedelta(days=80)),
        )
        adapter = ManualSnapshotAdapter({gdelt_rss_target(term):feed})
        app = create_app(db, proposal_model=model, start_worker=False, monitoring_adapters={'rss':adapter})
        client = TestClient(app, base_url='http://127.0.0.1')
        base = f'/api/projects/{project_id}'
        workspace = f'/api/workspace/projects/{project_id}'
        before = client.get(workspace+'/reading').json()
        thesis = client.get(workspace+'/profile').json()['fields']['thesis']
        check('baseline_and_update_saved', len(before['compiled']['batches']) == 2)

        # Actual DOCX upload and parser, not injected fragments or fake API responses.
        document = Document()
        document.add_paragraph('合成验收材料，不对应真实企业。')
        document.add_paragraph('财务：公司更正2025年未经审计实际收入为90万元；原说明书为100万元，差异原因未披露。')
        document.add_paragraph('客户：公司材料提到 Synthetic Peer 的产品试验，双方是否构成竞争关系尚未核实。')
        buffer = io.BytesIO()
        document.save(buffer)
        uploaded = client.post(base+'/files', files={
            'file':('合成经营更新.docx', buffer.getvalue(),
                    'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        check('docx_received', uploaded.status_code == 202)
        check('docx_parsed', app.state.intake_worker.process_one().status == 'completed')
        grant = app.state.reading_jobs.list_authorizations(project_id)[0]
        app.state.reading_jobs.set_status(grant['id'], 'paused', actor='synthetic-acceptance')
        check('pause_prevents_model_dispatch', app.state.reading_worker.process_one() is None and model.calls == 2)
        check('pause_preserves_existing_reading', client.get(workspace+'/reading').json()['compiled'] == before['compiled'])
        app.state.reading_jobs.set_status(grant['id'], 'active', actor='synthetic-acceptance')
        check('new_material_compiled', app.state.reading_worker.process_one()['state'] == 'succeeded')
        reading = client.get(workspace+'/reading').json()
        financial = next(u for u in reading['updates']['items'] if u['topic'] == 'financial')
        check('financial_versions_both_readable', '90万元' in financial['current'][0]['text']
              and '100万元' in financial['previous'][0]['text'])
        check('material_update_not_verified_business_event', financial['business_change_verified'] is False)
        check('human_thesis_unchanged', client.get(workspace+'/profile').json()['fields']['thesis'] == thesis)
        check('no_automatic_formal_confirmation', not repo.list_observations(project_id, status='confirmed'))
        for batch in reading['compiled']['batches']:
            for item in batch['items']:
                for ref in item['source_refs']:
                    source = repo.get_source(ref['source_id'])
                    check(f"source_quote_{batch['id']}_{item['topic']}_{ref['ordinal']}",
                          source.project_id == project_id and source.content_sha256 == ref['version']
                          and any(f.locator == ref['locator'] and ref['quote'] in f.text
                                  for f in repo.list_fragments(source.id)))
        calls = model.calls
        duplicate = client.post(base+'/files', files={'file':('合成经营更新.docx', buffer.getvalue(),
                                    'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        check('duplicate_no_model_call', duplicate.status_code == 202
              and app.state.reading_worker.process_one() is None and model.calls == calls)

        focus = base+'/monitoring/focus'
        selection = dict(category='competitors', term=term, provider='gdelt')
        preview = client.get(focus, params=selection).json()
        started = client.post(focus, json={**selection, 'consent_version':preview['consent_version']})
        check('public_scope_authorized', started.status_code == 202)
        check('public_job_finished', app.state.monitoring_scheduler.process_one().status == 'done')
        external = client.get(base+'/monitoring/external').json()
        stories = external['categories'][2]['items']
        check('old_and_other_entity_excluded', len(stories) == 1 and stories[0]['title'] == feed[0].title)
        context = stories[0]['internal_context']
        check('internal_external_link_has_quote_not_causal_claim', context['status'] == 'shared_term_only'
              and context['materials'] and context['impact'] is None and not context['relationship_verified'])
        check('headline_topic_not_verified_event', stories[0]['attention']['topic'] == 'performance'
              and not stories[0]['attention']['event_verified'])
        other = repo.create_project('另一家合成企业')
        foreign = client.get(f'/api/projects/{other.id}/monitoring/external').json()
        check('project_isolation', all(not g['items'] for g in foreign['categories']))

        # Receiver temporarily absent: keep reading, do not fall back to another service.
        missing_model = create_app(db, start_worker=False, monitoring_adapters={'rss':adapter})
        check('model_unavailable_keeps_reading', missing_model.state.reading_worker.process_one() is None
              and TestClient(missing_model).get(workspace+'/reading').json()['compiled'] == reading['compiled'])
        reboot = create_app(db, proposal_model=model, start_worker=False, monitoring_adapters={'rss':adapter})
        restarted = TestClient(reboot)
        check('restart_retains_internal_versions', restarted.get(workspace+'/reading').json()['compiled'] == reading['compiled'])
        check('restart_retains_external_context', restarted.get(base+'/monitoring/external').json() == external)
        check('restart_does_not_rebill', reboot.state.reading_worker.process_one() is None and model.calls == calls)
        plan_id = started.json()['plan']['id']
        # A later failure must not erase a previous successful result or say "no change".
        reboot.state.monitoring_scheduler.runner.adapters['rss'] = ManualSnapshotAdapter({})
        run = restarted.post(base+f'/monitoring/plans/{plan_id}/runs', json={'idempotency_key':'synthetic-failed-source'})
        check('retry_source_requested', run.status_code == 202)
        check('source_failure_processed', reboot.state.monitoring_scheduler.process_one().status == 'done')
        failed = restarted.get(base+'/monitoring/external').json()['categories'][2]
        check('source_failure_retains_last_news', failed['coverage'] == 'failed' and len(failed['items']) == 1)
        report['passed'] = all(report['checks'].values())
        report['synthetic_model_calls'] = calls
        return reboot, project_id, report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--serve', type=int, default=0)
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix='pw-monitor-journey-') as root:
        app, project_id, report = exercise(root)
        print(json.dumps(report, ensure_ascii=False), flush=True)
        if args.serve:
            import uvicorn
            print(f'Synthetic only: http://127.0.0.1:{args.serve}/#project/{project_id}/overview', flush=True)
            uvicorn.run(app, host='127.0.0.1', port=args.serve)


if __name__ == '__main__':
    main()
