"""Installed PPTX/XLSX intake-to-reading journey; synthetic, no network/model calls.

These are real Office containers parsed by the product, not injected fragments.
The deterministic compiler only exercises the integration contract, not AI quality.
"""
import argparse
import io
import json
import os
import re
import runpy
import tempfile
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from project_watch.api import create_app
from project_watch.repository import ProjectRepository

_fixture = runpy.run_path(str(Path(__file__).with_name('smoke_company_reading.py')))


class OfficeCompiler(_fixture['SyntheticCompiler']):
    """Known synthetic table convention only; never used by the application."""
    def compile_reading(self, packet, *, before_send):
        result = super().compile_reading(packet, before_send=before_send)
        if not packet['reference'].endswith('.xlsx'):
            return result
        rows = {f['locator']: f for f in packet['fragments']}
        scope, amount = rows['sheet 利润表 row 1'], rows['sheet 利润表 row 3']
        actual = re.search(r'B3\[header:2025年实际\]=(\d+)', amount['text']).group(1)
        budget = re.search(r'C3\[header:2026年预算\]=(\d+)', amount['text']).group(1)
        return {'items': [{'topic': 'financial', 'headline': '营业收入及预算口径',
            'text': f'公司材料列示2025年母公司营业收入{actual}万元，未经审计；2026年预算{budget}万元不是实际收入。',
            'as_of': '2025', 'evidence': [{'ref_id': f['ref_id'], 'quote': f['text']} for f in (scope, amount)]}]}


def bp_bytes():
    deck = Presentation()
    for lines in (
        ['合成光学：仅用于验收', '业务：合成光学公司提供工业检测设备，处于客户验证阶段。',
         '风险：客户试用不等于订单，需核实验收条件与收入确认依据。'],
        ['产品与客户', '客户：工业相机正在客户产线试用，尚未取得采购合同。'],
    ):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3)).text_frame.text = '\n'.join(lines)
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(7), Inches(1)).table
    for r, values in enumerate([['产品', '状态'], ['工业相机', '客户试用']]):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    stream = io.BytesIO()
    deck.save(stream)
    return stream.getvalue()


def finance_bytes(actual):
    book = Workbook()
    sheet = book.active
    sheet.title = '利润表'
    sheet.append(['合成光学公司母公司口径；币种：人民币；单位：万元；未经审计'])
    sheet.append(['项目', '2025年实际', '2026年预算'])
    sheet.append(['营业收入', actual, 120])
    sheet.append(['期末现金余额', None, None])
    note = book.create_sheet('说明')
    note.append(['合成材料；更正版' if actual == 90 else '合成材料；初版'])
    note.append(['现金余额未提供，不得视为零。'])
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def exercise(root):
    report = dict(passed=False, formats=['pptx', 'xlsx'], private_materials_used=False,
                  real_model_quality='not_tested', network_used=False, checks={})

    def check(name, value):
        report['checks'][name] = bool(value)
        assert value, name

    with patch.dict(os.environ, {'PROJECT_WATCH_SETTINGS_TOKEN': 'synthetic-local-only'}):
        model = OfficeCompiler()
        db = Path(root) / 'watch.sqlite3'
        app = create_app(db, proposal_model=model, start_worker=False)
        client = TestClient(app, base_url='http://127.0.0.1')
        project = client.post('/api/projects', json={'name': '合成光学 · Office 验收'}).json()
        project_id = project['id']
        base, workspace = f'/api/projects/{project_id}', f'/api/workspace/projects/{project_id}'
        repo = ProjectRepository(db)
        headers = {'Origin': 'http://127.0.0.1', 'X-Settings-Token': 'synthetic-local-only'}

        def upload(name, payload):
            mime = ('application/vnd.openxmlformats-officedocument.presentationml.presentation'
                    if name.endswith('.pptx') else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            received = client.post(base+'/files', files={'file': (name, payload, mime)})
            check('received_'+name, received.status_code == 202)
            processed = app.state.intake_worker.process_one()
            check('parsed_'+name, processed is not None and processed.status == 'completed')
            return received.json()['source']

        bp = upload('合成项目BP.pptx', bp_bytes())
        slides = client.get(f"/api/sources/{bp['id']}/fragments").json()
        check('pptx_slide_locators', [s['locator'] for s in slides] == ['slide 1', 'slide 2'])
        check('pptx_table_text_preserved', '工业相机 | 客户试用' in slides[1]['text'])
        check('no_call_before_consent', app.state.reading_worker.process_one() is None and model.calls == 0)
        disclosure = client.get(workspace+'/reading-authorization').json()
        grant = client.post(workspace+'/reading-authorization', headers=headers, json={
            'consent': True, 'config_id': disclosure['model']['config_id'], 'request_key': 'synthetic-office',
            'rule_version': disclosure['rule_version'], 'max_output_tokens': disclosure['model']['max_output_tokens'],
            'include_future': True, 'max_requests': 3})
        check('consent_saved', grant.status_code == 200)
        check('pptx_compiled', app.state.reading_worker.process_one()['state'] == 'succeeded')
        profile = client.get(workspace+'/profile').json()
        check('bp_populates_profile', '工业检测设备' in profile['fields']['business']['text']
              and profile['focus'][0]['source_refs'][0]['locator'] == 'slide 1')
        repo.edit_monitor_profile(project_id, expected_version=0, fields={'thesis': '合成验收：尚不能据此确认收入质量。'})
        thesis = client.get(workspace+'/profile').json()['fields']['thesis']

        first = upload('合成财务初版.xlsx', finance_bytes(100))
        check('xlsx_compiled', app.state.reading_worker.process_one()['state'] == 'succeeded')
        revised_bytes = finance_bytes(90)
        revised = upload('合成财务更正版.xlsx', revised_bytes)
        check('xlsx_update_compiled', app.state.reading_worker.process_one()['state'] == 'succeeded')
        reading = client.get(workspace+'/reading').json()
        change = next(u for u in reading['updates']['items'] if u['topic'] == 'financial')
        check('versions_both_readable', '90万元' in change['current'][0]['text'] and '100万元' in change['previous'][0]['text'])
        check('period_scope_budget_qualifiers', all(
            all(word in item['text'] for word in ('2025年', '母公司', '万元', '未经审计', '2026年预算120万元不是实际'))
            for item in (*change['current'], *change['previous'])))
        check('unreviewed_comparison_not_formal_event', not change['business_change_verified'])
        check('file_versions_distinct', first['id'] != revised['id'] and first['content_sha256'] != revised['content_sha256'])
        cells = client.get(f"/api/sources/{revised['id']}/fragments").json()
        check('xlsx_cell_header_and_sheet_preserved', any(
            f['locator'] == 'sheet 利润表 row 3' and 'B3[header:2025年实际]=90' in f['text']
            and 'C3[header:2026年预算]=120' in f['text'] for f in cells))
        check('missing_cell_not_zero', next(f['text'] for f in cells if f['locator'] == 'sheet 利润表 row 4') == 'A4=期末现金余额')
        for batch in reading['compiled']['batches']:
            for item in batch['items']:
                for ref in item['source_refs']:
                    source = repo.get_source(ref['source_id'])
                    check(f"citation_{batch['id']}_{item['topic']}_{ref['ordinal']}",
                          source.project_id == project_id and source.content_sha256 == ref['version']
                          and any(f.locator == ref['locator'] and ref['quote'] in f.text for f in repo.list_fragments(source.id)))
        duplicate = client.post(base+'/files', files={'file': ('合成财务更正版.xlsx', revised_bytes,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        check('duplicate_not_billed', duplicate.status_code == 202 and app.state.reading_worker.process_one() is None and model.calls == 3)
        reboot = create_app(db, proposal_model=model, start_worker=False)
        check('restart_keeps_reading', TestClient(reboot).get(workspace+'/reading').json()['compiled'] == reading['compiled'])
        check('restart_not_billed', reboot.state.reading_worker.process_one() is None and model.calls == 3)
        check('human_thesis_retained', TestClient(reboot).get(workspace+'/profile').json()['fields']['thesis'] == thesis)
        check('no_automatic_confirmed_fact', not repo.list_observations(project_id, status='confirmed'))
        report['passed'] = all(report['checks'].values())
        report['synthetic_model_calls'] = model.calls
        return reboot, project_id, report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--serve', type=int, default=0)
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix='pw-office-reading-') as root:
        app, project_id, report = exercise(root)
        # Windows CI may use cp1252 stdout. JSON escapes preserve Chinese names
        # losslessly without requiring a particular operator console encoding.
        print(json.dumps(report, ensure_ascii=True), flush=True)
        if args.serve:
            import uvicorn
            print(f'Synthetic only: http://127.0.0.1:{args.serve}/#project/{project_id}/internal', flush=True)
            uvicorn.run(app, host='127.0.0.1', port=args.serve)


if __name__ == '__main__':
    main()
