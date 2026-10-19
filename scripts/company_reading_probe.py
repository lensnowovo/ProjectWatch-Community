"""Bounded DSH long-document acceptance. Preview by default; built-in synthetic data only.

Run with the installed interpreter (-I) to test the installed package. No pilot DB,
input file or service activation option exists. --run requires operator approval
outside this script. Keys enter through stdin or existing process environment only.
Integration checks do NOT certify semantic quality: inspect the returned readings.
"""

import argparse
import hashlib
import io
import json
import os
import secrets
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from pydantic import ValidationError

from project_watch.api import create_app
from project_watch.company_state import (
    READING_OUTPUT_TOKENS,
    ReadingOutput,
    input_packets,
    validate_reading,
)
from project_watch.dsh_executor import DshProposalModel
from project_watch.model_provider import ModelProviderConfig
from project_watch.repository import ProjectRepository

BASELINE = """合成光学公司说明书。以下全部是验收专用虚构信息，不对应任何真实公司。
业务：公司提供工业视觉检测设备，服务零部件工厂，处于客户验证阶段。
公司的设备用于发现加工件表面缺陷，输出图像及疑似缺陷位置，由工厂人员复核。设备不能替代客户最终质量责任，说明书未提供事故责任分担条款。公司没有提供独立第三方检测结论，因此资料中的精度描述仅属于自述，不能写成已经验证的性能。产品包括图像采集模组、检测软件和产线集成服务。标准硬件需要适配客户工位及光照条件，软件需要使用经授权的样本进行调试。现场环境改变可能需要重新标定，交付人员尚未证明不同工厂之间能够直接复制同一配置。
客户：设备正在客户产线试用，验收及采购合同尚未取得。
试用反馈由客户工艺人员提供，公司暂未取得采购部门的书面订单。试用项目名称不构成客户背书，公司展示的客户标识不代表已建立持续供货关系。研究时需要分别核对测试邀请、试用协议、验收记录和采购合同，不能把几个阶段合并为已经完成商业化。
公司表示拟采用设备销售与后续维护相结合的商业模式，但未披露维护费收取周期、服务范围和续费记录。当前收入是否包括现场开发服务、硬件转售或关联交易尚不明确。报价单不能替代实际成交价格，尚无法据此判断毛利水平。研发主要围绕成像稳定性、缺陷样本管理和现场部署展开。项目资料没有列明专利授权号、权利人或有效状态，不应据此评价知识产权壁垒。是否依赖外部算法授权、客户样本能否跨项目复用，以及研发人员的职务成果归属，需要另行取得合同证据。
公司将照明组件与相机采购交由外部供应商，内部负责装配、软件部署和交付。资料没有提供供应商集中度、交货周期和备选供应商认证情况。采购品牌可用不等于供应链风险已经解决，不能仅根据产品清单推断替换成本或议价能力。目前产能描述采用理论装配能力口径，没有提供实际工时、良率、产线利用率和售后返工记录。试生产批次与正式交付批次也没有明确区分。对于扩产计划，应先确认订单与交付要求，再评估设备投入，不能把计划产能作为已经实现的销售能力。
交付流程包括现场勘查、样本确认、设备调试、试运行与客户验收。每个阶段可能由不同客户人员负责，但资料未说明签字权限。合同付款条件尚未披露，客户试用结束与收入确认之间的关系也未明确，不宜根据交付照片确认收入。售后支持以现场调整和远程诊断为主，服务响应时间属于公司提出的目标，尚无实际履约统计。资料没有提供返修率、质保费用或重大质量争议清单。投研人员需要获取历史服务记录，区分偶发调试问题和可能持续发生的质量问题。
管理团队的职责划分只出现在组织图中，没有附劳动合同、履历证明或股权激励文件。不能依据职务名称推断人员已全职到岗。关键研发人员是否承担其他公司的工作、是否存在竞业约束，以及核心客户关系由谁维护，目前没有足够依据。公司拟优先拓展已完成试用的行业客户，尚未提供可验证的市场规模研究。材料中提及的行业增长只是定性背景，不能直接转换为本公司订单增长。市场需求、设备验证周期与客户资本开支计划之间的关系仍需逐项核实。
竞争方面，公司自称具备现场部署经验，但没有提供同条件对比测试。材料没有出现经核实的竞品名称，不应自行添加企业作为直接竞争对手。后续比较需要保持相同应用场景、检测对象、速度和误判口径，不能只比较宣传页参数。
经营记录没有说明收入确认政策、退货安排或是否存在渠道压货。应收款按账面余额列示，尚未取得客户对账和期后回款证明。收入规模、现金回款和订单金额属于不同指标，不能互相代替，更不能将应收账款余额当作新增收入。
融资资料没有提供本轮估值、已签投资协议或交割条件。管理层拟用于研发与客户交付的资金安排仍是计划，不表示已取得融资。股东借款、历史融资和经营现金流没有在同一张表中列明，现有资料不足以计算现金可支撑的月份。
本说明书提供的是公司自述及资料缺口，不构成审计、法律或投资意见。研究人员应保留独立判断。若新文件与本说明书矛盾，应并列展示期间、主体与出处，而不是直接选择更新日期更近的数字。未知项目保持未知，不能填零。
财务：公司材料列示2025年收入100万元，属于未经审计的实际数，回款尚未核实。
财务：公司预计2026年收入180万元，属于预测，不是已实现收入；未提供预测订单明细。
风险：客户试用不等于订单，应核实验收条件与收入确认依据，暂不形成投资结论。
"""

UPDATE = """合成光学经营更新，全部为虚构验收信息。
客户：新增客户完成试用，但尚未签署采购合同；原有试用客户仍待验收。
财务：公司更正2025年未经审计实际收入为90万元，原说明书列示100万元；差异原因未披露，需核对明细。
财务：2026年预测收入仍为180万元，不能作为实际收入，也不能与2025年实际数直接计算已实现增长。
风险：经营更新没有审计报告或回款凭证，收入更正不能自动替代人工结论。
"""

DOCUMENTS = (("合成公司说明书.txt", BASELINE), ("合成经营更新.txt", UPDATE))
HUMAN_THESIS = "先核实客户验收与实际回款，暂不采纳公司预测。"


def document_names(document_format):
    if document_format not in {'text', 'office'}:
        raise ValueError('unsupported built-in corpus format')
    extensions = ('pptx', 'docx') if document_format == 'office' else ('txt', 'txt')
    return [str(Path(name).with_suffix('.' + extension))
            for (name, _), extension in zip(DOCUMENTS, extensions, strict=True)]


def document_payloads(document_format):
    """Serialize only the fixed synthetic corpus; never accept operator file paths."""
    names = document_names(document_format)
    if document_format == 'text':
        return [(name, text.encode(), 'text/plain')
                for name, (_, text) in zip(names, DOCUMENTS, strict=True)]
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    lines = BASELINE.splitlines()
    for start in range(0, len(lines), 4):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5)).text_frame.text = '\n'.join(lines[start:start + 4])
    bp = io.BytesIO()
    deck.save(bp)
    doc = Document()
    for line in UPDATE.splitlines():
        doc.add_paragraph(line)
    update = io.BytesIO()
    doc.save(update)
    prefix = 'application/vnd.openxmlformats-officedocument.'
    return [(names[0], bp.getvalue(), prefix + 'presentationml.presentation'),
            (names[1], update.getvalue(), prefix + 'wordprocessingml.document')]


def manifest(document_format='text'):
    names = document_names(document_format)
    return {
        "mode": "preview_no_network", "corpus": "built_in_synthetic_only",
        "formats": ['pptx', 'docx'] if document_format == 'office' else ['txt', 'txt'],
        "max_requests": 2, "max_output_tokens_per_request": READING_OUTPUT_TOKENS,
        "automatic_retry": False, "pilot_database_access": False,
        "document_hash_scope": "utf8_source_text; actual containers are hashed separately during execution",
        "documents": [{"name": name, "chars": len(text),
                       "sha256": hashlib.sha256(text.encode()).hexdigest()}
                      for name, (_, text) in zip(names, DOCUMENTS, strict=True)],
    }


def load_model(stream=None):
    if stream is None:
        config = ModelProviderConfig.from_environment(os.environ)
        key = config.resolve_key(os.environ)
    else:
        raw = stream.read(16385)
        if len(raw) > 16384:
            raise ValueError("configuration too large")
        fields = json.loads(raw)
        required = {"base_url", "protocol", "model", "api_key"}
        if (not isinstance(fields, dict) or not required <= set(fields) <= required | {"thinking_mode"}
                or any(not isinstance(value, str) for value in fields.values())):
            raise ValueError("invalid configuration")
        key = fields.pop("api_key")
        config = ModelProviderConfig(**fields)
        config.resolve_key({config.credential_env: key})
    return DshProposalModel(config, key, max_runs=2)


class ProbeFailed(Exception):
    pass


def diagnose_output(packet, output):
    """Synthetic probe only: no provider exception bodies or arbitrary fields."""
    try:
        candidate = ReadingOutput.model_validate(output).model_dump()
    except ValidationError as exc:
        known_fields = {'items', 'topic', 'headline', 'text', 'evidence', 'as_of', 'ref_id', 'quote'}
        errors = exc.errors(include_input=False, include_context=False)
        return {'code': 'schema_invalid', 'error_types': sorted({
            error['type'] for error in errors
        }), 'locations': [[part if type(part) is int or part in known_fields else '<unknown>'
                          for part in error['loc']] for error in errors[:30]]}
    try:
        validate_reading(packet, candidate)
    except ValueError as exc:
        known = {
            'reading output too large', 'empty reading',
            'reading quote not present in authorized packet',
            'reading number not present in supporting quote',
            'reading period not present in supporting quote',
        }
        return {'code': str(exc) if str(exc) in known else 'validation_failed',
                'candidate': candidate}
    return {'code': 'passed_quote_checks_not_semantics'}


def exercise(model, *, document_format='text'):
    """Real intake/consent/reading APIs with no background scheduler or live service."""
    documents = document_payloads(document_format)
    result = {**manifest(document_format), "mode": "execution", "checks": {}, "turns": [],
              "semantic_review": "required", "passed": False}
    result['container_hashes'] = [hashlib.sha256(data).hexdigest() for _, data, _ in documents]
    compile_reading = model.compile_reading
    diagnostics = []

    def observe(packet, *, before_send):
        output = compile_reading(packet, before_send=before_send)
        diagnostics.append(diagnose_output(packet, output))
        return output

    def check(name, condition):
        result["checks"][name] = bool(condition)
        if not condition:
            raise ProbeFailed(name)

    def capture_turn(label, job):
        # Allowlist numeric usage only, never provider messages or errors.
        receipt = getattr(model, "receipt", {})
        usage = {key: value for key in ("requests_admitted", "observed_prompt_tokens", "observed_completion_tokens")
                 if type(value := receipt.get(key)) is int and value >= 0}
        result["turns"].append({"stage": label, "state": job["state"] if job else None, "usage": usage,
                                "validation": diagnostics.pop() if diagnostics else {
                                    "code": "no_structured_response"}})

    token = secrets.token_urlsafe(32)
    with tempfile.TemporaryDirectory(prefix="pw-live-synthetic-") as root, patch.dict(
        os.environ, {"PROJECT_WATCH_SETTINGS_TOKEN": token}
    ), patch.object(model, "compile_reading", observe):
        db = Path(root) / "synthetic.sqlite3"
        app = create_app(db, proposal_model=model, start_worker=False)
        repo = ProjectRepository(db)
        try:
            with TestClient(app, base_url="http://127.0.0.1") as client:
                project = client.post("/api/projects", json={"name": "合成光学 · DSH验收"}).json()
                project_id = project["id"]
                route = f"/api/workspace/projects/{project_id}"
                upload = f"/api/projects/{project_id}/files"

                def add_document(index):
                    name, data, mime = documents[index]
                    response = client.post(upload, files={"file": (name, data, mime)})
                    check(f"upload_{index}", response.status_code == 202)
                    parsed = app.state.intake_worker.process_one()
                    check(f"parse_{index}", parsed is not None and parsed.status == "completed")
                    source = next(s for s in repo.list_sources(project_id)
                                  if s.content_sha256 == hashlib.sha256(data).hexdigest())
                    packets = input_packets(repo, source)
                    check(f"single_packet_{index}", len(packets) == 1)
                    fragments = packets[0]['fragments']
                    parsed_text = '\n'.join(f['text'] for f in fragments)
                    check(f'all_text_preserved_{index}', all(
                        line in parsed_text for line in DOCUMENTS[index][1].splitlines()))
                    if document_format == 'office':
                        prefix = 'slide ' if index == 0 else 'paragraph '
                        check(f'office_locations_{index}', all(
                            f['locator'].startswith(prefix) for f in fragments))
                    if index == 0:
                        check("long_document_not_excerpt", sum(len(f["text"]) for f in packets[0]["fragments"]) > 1500)

                add_document(0)
                check("no_call_without_consent", app.state.reading_worker.process_one() is None)
                disclosure = client.get(route + "/reading-authorization").json()
                grant = client.post(route + "/reading-authorization", headers={
                    "Origin": "http://127.0.0.1", "X-Settings-Token": token,
                }, json={"consent": True, "config_id": disclosure["model"]["config_id"],
                         "rule_version": disclosure["rule_version"],
                         "max_output_tokens": disclosure["model"]["max_output_tokens"],
                         "request_key": "synthetic-long-reading-probe", "include_future": True,
                         "max_requests": 2})
                check("consent_saved", grant.status_code == 200)
                job = app.state.reading_worker.process_one()
                capture_turn("baseline", job)
                check("baseline_published", job is not None and job["state"] == "succeeded")
                baseline = client.get(route + "/reading").json()["compiled"]
                check("baseline_nonempty", len(baseline["batches"]) == 1 and bool(baseline["batches"][0]["items"]))
                # Check idempotency while there is still unused quota. A no-op
                # after exhausting quota would not prove duplicate prevention.
                name, data, mime = documents[0]
                check("duplicate_accepted", client.post(upload, files={"file": (name, data, mime)}).status_code == 202)
                check("duplicate_not_billed", app.state.reading_worker.process_one() is None)
                check("unused_budget_after_duplicate", app.state.reading_jobs.status(grant.json()["id"])["requests_used"] == 1)
                reboot = create_app(db, proposal_model=model, start_worker=False)
                check("no_dispatch_after_restart", reboot.state.reading_worker.process_one() is None)
                repo.edit_monitor_profile(project_id, expected_version=0, fields={"thesis": HUMAN_THESIS})
                add_document(1)
                job = app.state.reading_worker.process_one()
                capture_turn("update", job)
                check("update_published", job is not None and job["state"] == "succeeded")
                result["reading"] = client.get(route + "/reading").json()["compiled"]
                check("both_versions_retained", len(result["reading"]["batches"]) == 2)
                check("new_reading_nonempty", all(b["items"] for b in result["reading"]["batches"]))
                thesis = client.get(route + "/profile").json()["fields"]["thesis"]
                check("human_judgment_preserved", thesis["origin"] == "human" and thesis["text"] == HUMAN_THESIS)
                grant_status = app.state.reading_jobs.status(grant.json()["id"])
                check("two_dispatches_only", grant_status["requests_used"] == 2)
                check("no_formal_facts", repo.list_observations(project_id, status="confirmed") == [])
                # Recreate repository/worker against the same persisted synthetic DB.
                reboot = create_app(db, proposal_model=model, start_worker=False)
                check("budget_exhausted", reboot.state.reading_worker.status(project_id)["reason"] == "budget_reached")
                check("no_dispatch_after_budget", reboot.state.reading_worker.process_one() is None)
                with TestClient(reboot, base_url="http://127.0.0.1") as restarted:
                    check("reading_survives_restart", restarted.get(route + "/reading").json()["compiled"] == result["reading"])
                result["passed"] = True
        except ProbeFailed as exc:
            result["failed_check"] = str(exc)
        finally:
            app.state.reading_worker.stop()
    return result


def main(argv=None, *, stream=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", action="store_true", help="explicitly run up to two model calls")
    parser.add_argument("--config-stdin", action="store_true", help="receive provider configuration through stdin")
    parser.add_argument('--format', choices=('text', 'office'), default='text',
                        help='built-in TXT pair or actual PPTX/DOCX containers; same two-call limit')
    args = parser.parse_args(argv)
    if not args.run:
        result = manifest(args.format)
    else:
        try:
            model = load_model((stream or sys.stdin) if args.config_stdin else None)
            result = exercise(model, document_format=args.format)
        except Exception:
            result = {"passed": False, "error": "Acceptance could not finish; configuration and provider details suppressed."}
    # Portable JSON receipt: Windows redirected stdout may be cp1252 even with
    # -I. Escaping preserves every Unicode character without a console crash.
    print(json.dumps(result, ensure_ascii=True, indent=2))
    return 2 if result.get("passed") is False else 0


if __name__ == "__main__":
    raise SystemExit(main())
